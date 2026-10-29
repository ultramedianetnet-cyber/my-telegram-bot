import asyncio
import html
import json
import logging
import os
import re
import sqlite3
import time
from io import BytesIO

from dotenv import load_dotenv

load_dotenv()

from aiogram import BaseMiddleware, Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError
from aiogram.filters import BaseFilter, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    CallbackQuery,
    ErrorEvent,
    BufferedInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputMediaPhoto,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
)
from aiogram.utils.media_group import MediaGroupBuilder
import db
import db

# ----------------------------------------------------------------- SOZLAMALAR
TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_IDS = [int(x) for x in os.getenv("ADMIN_IDS", "").replace(" ", "").split(",") if x]
VIDEO_GROUP_ID = int(os.getenv("VIDEO_GROUP_ID", "0") or 0)
ABOUT = os.getenv("ABOUT_TEXT", "Matematika o'quv markazi.\nAloqa va kanal ma'lumotlari tez orada.")
SUBJECTS = ["Algebra", "Geometriya", "Boshqa"]

if not TOKEN:
    raise SystemExit("BOT_TOKEN topilmadi. .env faylini to'ldiring.")

logging.basicConfig(level=logging.INFO)
bot = Bot(TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
dp = Dispatcher(storage=MemoryStorage())
r = Router()
esc = html.escape

# --------------------------------------------------------------- TUGMA MATNLARI
B_TEST = "📝 MOCK test"
B_VID = "🎬 Video darslar"
B_BOOK = "📚 Kerakli kitoblar"
B_ABOUT = "ℹ️ Biz haqimizda"
B_ADMIN = "🛠 Admin panel"
B_BACK = "⬅️ Asosiy menyu"
A_ADDST = "➕ O'quvchi qo'shish"
A_STUDENTS = "👨‍🎓 O'quvchilar"
A_GROUPS = "👥 Guruhlar"
A_VADD = "🎬 Video qo'shish"
A_VATT = "📎 Videoni guruhga biriktirish"
A_VDEL = "🗑 Videoni o'chirish"
A_RES = "🧾 Test natijalari"
A_TADD = "➕ Test qo'shish"
A_TEDIT = "✏️ Testni tahrirlash"
A_TDEL = "🗑 Testni o'chirish"
A_BADD = "📖 Kitob qo'shish"
A_BEDIT = "✏️ Kitobni tahrirlash"
A_BDEL = "🗑 Kitobni o'chirish"
A_BC = "📢 Xabar yuborish"
A_STAT = "📊 Statistika"
A_REPORT = "📑 Hisobot"
A_ADMINS = "👮 Adminlar"


def rkb(rows):
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text=t) for t in row] for row in rows],
        resize_keyboard=True,
    )


def ikb(rows):
    return InlineKeyboardMarkup(
        inline_keyboard=[[InlineKeyboardButton(text=t, callback_data=d) for t, d in row] for row in rows]
    )


def main_kb(admin):
    rows = [[B_TEST, B_VID], [B_BOOK, B_ABOUT]]
    if admin:
        rows.append([B_ADMIN])
    return rkb(rows)


def admin_kb():
    return rkb([
        [A_ADDST, A_STUDENTS],
        [A_GROUPS],
        [A_VADD, A_VATT, A_VDEL],
        [A_RES, A_REPORT],
        [A_TADD, A_TEDIT, A_TDEL],
        [A_BADD, A_BEDIT, A_BDEL],
        [A_BC, A_STAT],
        [A_ADMINS, B_BACK],
    ])


# ------------------------------------------------------------------- YORDAMCHI
async def is_admin(uid):
    return await db.one("SELECT 1 FROM admins WHERE id=?", (uid,)) is not None


class IsAdmin(BaseFilter):
    async def __call__(self, event) -> bool:
        return event.from_user is not None and await is_admin(event.from_user.id)


async def menu_for(uid):
    return main_kb(await is_admin(uid))


async def safe_send(uid, text, markup=None):
    try:
        await bot.send_message(uid, text, reply_markup=markup)
        return True
    except TelegramForbiddenError:
        await db.run("UPDATE users SET blocked=1 WHERE id=?", (uid,))
    except Exception as e:
        logging.warning("send fail %s: %s", uid, e)
    return False


async def get_student(obj):
    """Ro'yxatdan o'tgan va guruhga qo'shilgan o'quvchini qaytaradi."""
    msg = obj if isinstance(obj, Message) else obj.message
    u = await db.one("SELECT * FROM users WHERE id=?", (obj.from_user.id,))
    if not u:
        await msg.answer("Avval /start bosib ro'yxatdan o'ting.")
        return None
    if not u["group_id"]:
        await msg.answer("⏳ Siz hali guruhga qo'shilmagansiz. Admin sizni guruhga qo'shishini kuting.")
        return None
    return u


async def send_content(chat_id, ctype, file_id, text, caption=None, markup=None):
    file_caption = "\n\n".join(
        part for part in (
            f"<b>{esc(caption)}</b>" if caption else "",
            esc(text or ""),
        ) if part
    )
    if ctype == "photo":
        await bot.send_photo(chat_id, file_id, caption=file_caption or None, reply_markup=markup)
    elif ctype == "document":
        await bot.send_document(chat_id, file_id, caption=file_caption or None, reply_markup=markup)
    else:
        body = f"<b>{esc(caption)}</b>\n\n{esc(text or '')}" if caption else esc(text or "")
        await bot.send_message(chat_id, body, reply_markup=markup)


# ---------------------------------------------------------------------- HOLATLAR
class Reg(StatesGroup):
    name = State()
    phone = State()


class GrpNew(StatesGroup):
    name = State()


class GrpEdit(StatesGroup):
    name = State()


class StudentEdit(StatesGroup):
    name = State()
    phone = State()


class TestAdd(StatesGroup):
    title = State()
    content = State()
    code = State()
    confirm_code = State()
    mode = State()
    question_count = State()
    auto_answer = State()


class TestAccess(StatesGroup):
    code = State()


class VideoAdd(StatesGroup):
    upload = State()


class TestEdit(StatesGroup):
    title = State()
    code = State()
    manual_content = State()
    auto_count = State()
    auto_answer = State()


class BookAdd(StatesGroup):
    file = State()


class BookEdit(StatesGroup):
    details = State()
    file = State()


class Submit(StatesGroup):
    photos = State()


class Review(StatesGroup):
    comment = State()


class Bc(StatesGroup):
    msg = State()


class AdmAdd(StatesGroup):
    name = State()
    uid = State()


class AdmEdit(StatesGroup):
    name = State()


# --------------------------------------------------------------- MIDDLEWARE
class Activity(BaseMiddleware):
    async def describe_callback(self, action):
        parts = action.split(":")
        if parts[0] == "vs" and len(parts) == 2:
            return f"Video darslar fani tanlandi: {parts[1]}"
        if parts[0] == "vl" and len(parts) == 2 and parts[1].isdigit():
            video = await db.one("SELECT subject,title FROM videos WHERE id=?", (int(parts[1]),))
            if video:
                return f"Video ko'rildi: {video['subject']} — {video['title']}"
        if parts[0] == "tp" and len(parts) == 2 and parts[1].isdigit():
            test_id = int(parts[1])
        elif parts[0] == "quiz" and len(parts) == 3 and parts[1] == "start" and parts[2].isdigit():
            test_id = int(parts[2])
        else:
            test_id = None
        if test_id is not None:
            test = await db.one("SELECT title FROM tests WHERE id=?", (test_id,))
            if test:
                return f"Test ochildi: {test['title']}"
        if parts[0] == "quiz" and len(parts) == 3 and parts[1] == "resume" and parts[2].isdigit():
            test = await db.one(
                "SELECT t.title FROM attempts a JOIN tests t ON t.id=a.test_id WHERE a.id=?",
                (int(parts[2]),),
            )
            if test:
                return f"Test davom ettirildi: {test['title']}"
        if parts[0] in {"qa", "qf"} and len(parts) >= 2 and parts[1].isdigit():
            attempt = await db.one(
                "SELECT t.title FROM attempts a JOIN tests t ON t.id=a.test_id WHERE a.id=?",
                (int(parts[1]),),
            )
            if attempt and parts[0] == "qa" and len(parts) == 4 and parts[2].isdigit():
                return f"{attempt['title']} — {int(parts[2]) + 1}-savol javobi: {parts[3]}"
            if attempt and parts[0] == "qf":
                return f"Test yakunlandi: {attempt['title']}"
        if parts[0] == "tsub" and len(parts) == 2 and parts[1].isdigit():
            test = await db.one("SELECT title FROM tests WHERE id=?", (int(parts[1]),))
            if test:
                return f"Qo'lda tekshiriladigan test topshirildi: {test['title']}"
        return action

    async def __call__(self, handler, event, data):
        user = data.get("event_from_user")
        if user:
            now = int(time.time())
            await db.run(
                "UPDATE users SET last_active=?, username=? WHERE id=?",
                (now, user.username, user.id),
            )
            if await db.one("SELECT 1 FROM users WHERE id=?", (user.id,)) and not await is_admin(user.id):
                if isinstance(event, CallbackQuery):
                    action = event.data or "button"
                    event_type = "callback"
                    description = await self.describe_callback(action)
                elif isinstance(event, Message):
                    action = event.text if event.text in {B_TEST, B_VID} else (
                        "text message" if event.text else
                        "media with caption" if event.caption else
                        "contact shared" if event.contact else "media sent"
                    )
                    event_type = "message"
                    description = {
                        B_TEST: "MOCK test menyusi ochildi",
                        B_VID: "Video darslar menyusi ochildi",
                    }.get(action, action)
                else:
                    action, event_type = "activity", "other"
                    description = action
                action = str(action).strip()
                if re.fullmatch(r"\d{4}", action):
                    action = "4-digit test code entered"
                category = "other"
                if action == B_VID or action.startswith(("vs:", "vl:", "vpage:")):
                    category = "video"
                elif action == B_TEST or action.startswith((
                    "tp:", "quiz:", "q:", "qa:", "qr:", "qe:", "qf:", "tsub:", "tests:legacy", "tdone"
                )):
                    category = "test"
                await db.run(
                    "INSERT INTO user_activity(user_id,action,description,category,event_type,occurred_at) "
                    "VALUES(?,?,?,?,?,?)",
                    (user.id, action[:160], str(description)[:250], category, event_type, now),
                )
        return await handler(event, data)


# ================================================================== RO'YXATDAN O'TISH
@r.message(CommandStart())
async def start(m: Message, state: FSMContext):
    await state.clear()
    u = await db.one("SELECT 1 FROM users WHERE id=?", (m.from_user.id,))
    if u or await is_admin(m.from_user.id):
        await m.answer("Asosiy menyu 👇", reply_markup=await menu_for(m.from_user.id))
        return
    await m.answer(
        "Assalomu alaykum! 👋\nMatematika o'quv botiga xush kelibsiz.\n\n"
        "Ism va familiyangizni yozing:",
        reply_markup=ReplyKeyboardRemove(),
    )
    await state.set_state(Reg.name)


# ---- menyu tugmalari (har qanday holatda ishlaydi)
@r.message(F.text == B_BACK)
async def back(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Asosiy menyu 👇", reply_markup=await menu_for(m.from_user.id))


@r.message(F.text == B_ABOUT)
async def about(m: Message, state: FSMContext):
    await state.clear()
    await m.answer(ABOUT)


@r.message(F.text == B_ADMIN, IsAdmin())
async def admin_panel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("🛠 Admin panel", reply_markup=admin_kb())


@r.message(Reg.name, F.text)
async def reg_name(m: Message, state: FSMContext):
    name = m.text.strip()
    if len(name) < 3:
        return await m.answer("Iltimos, ism va familiyani to'liq yozing.")
    await state.update_data(name=name)
    await state.set_state(Reg.phone)
    await m.answer(
        "Telefon raqamingizni yuboring (pastdagi tugmani bosing):",
        reply_markup=ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text="📱 Raqamni yuborish", request_contact=True)]],
            resize_keyboard=True,
            one_time_keyboard=True,
        ),
    )


@r.message(Reg.phone, F.contact)
async def reg_phone(m: Message, state: FSMContext):
    if m.contact.user_id != m.from_user.id:
        return await m.answer("Iltimos, o'zingizning raqamingizni yuboring.")
    d = await state.get_data()
    now = int(time.time())
    await db.run(
        "INSERT OR REPLACE INTO users(id,full_name,phone,username,created_at,last_active) VALUES(?,?,?,?,?,?)",
        (m.from_user.id, d["name"], m.contact.phone_number, m.from_user.username, now, now),
    )
    await db.run(
        "INSERT INTO user_activity(user_id,action,description,category,event_type,occurred_at) "
        "VALUES(?,?,?,?,?,?)",
        (m.from_user.id, "registration", "Ro'yxatdan o'tdi", "account", "system", now),
    )
    await state.clear()
    await m.answer(
        "✅ Ro'yxatdan o'tdingiz!\nAdmin sizni guruhga qo'shgach, darslar ochiladi.",
        reply_markup=await menu_for(m.from_user.id),
    )
    await send_app_button(m)
    for a in await db.fetch("SELECT id FROM admins"):
        await safe_send(a["id"], f"🆕 Yangi foydalanuvchi: <b>{esc(d['name'])}</b>\nUni «O'quvchi qo'shish» orqali guruhga qo'shing.")


@r.message(Reg.phone)
async def reg_phone_bad(m: Message):
    await m.answer("Pastdagi «📱 Raqamni yuborish» tugmasini bosing.")


# ================================================================== O'QUVCHI: VIDEO
@r.message(F.text == B_VID)
async def vid_menu(m: Message, state: FSMContext):
    await state.clear()
    u = await get_student(m)
    if not u:
        return
    g = await db.one("SELECT name FROM groups WHERE id=?", (u["group_id"],))
    await m.answer(
        f"👥 Guruh: <b>{esc(g['name'])}</b>\nFanni tanlang:",
        reply_markup=ikb([[("📐 Algebra", "vs:Algebra"), ("📏 Geometriya", "vs:Geometriya")], [("📂 Boshqa", "vs:Boshqa")]]),
    )


@r.callback_query(F.data.startswith("vs:"))
async def vid_subject(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    subj = c.data.split(":", 1)[1]
    markup = await video_list_kb(u["group_id"], subj, 0)
    if markup is None:
        return await c.answer("Hozircha bu fandan darslar yo'q.", show_alert=True)
    await c.message.answer(
        f"<b>{subj}</b> — mavzular:",
        reply_markup=markup,
    )
    await c.answer()


async def video_list_kb(group_id, subject, page):
    total = await db.count(
        "SELECT COUNT(*) FROM videos v JOIN group_videos gv ON gv.video_id=v.id "
        "WHERE gv.group_id=? AND v.subject=?",
        (group_id, subject),
    )
    if not total:
        return None
    rows = await db.fetch(
        "SELECT v.id, v.title FROM videos v JOIN group_videos gv ON gv.video_id=v.id "
        "WHERE gv.group_id=? AND v.subject=? ORDER BY gv.position, v.id LIMIT ? OFFSET ?",
        (group_id, subject, VIDEO_PAGE_SIZE, page * VIDEO_PAGE_SIZE),
    )
    start = page * VIDEO_PAGE_SIZE
    kb = [
        [(f"{start + index + 1}-dars: {row['title']}"[:60], f"vl:{row['id']}")]
        for index, row in enumerate(rows)
    ]
    nav = []
    if page > 0:
        nav.append(("⬅️ Oldingi", f"vpage:{subject}:{page - 1}"))
    if start + VIDEO_PAGE_SIZE < total:
        nav.append(("Keyingi ➡️", f"vpage:{subject}:{page + 1}"))
    if nav:
        kb.append(nav)
    return ikb(kb)


@r.callback_query(F.data.startswith("vpage:"))
async def video_list_page(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    _, subject, page = c.data.split(":")
    if subject not in SUBJECTS or not page.isdigit():
        return await c.answer("Sahifa topilmadi.", show_alert=True)
    markup = await video_list_kb(u["group_id"], subject, int(page))
    if markup is None:
        return await c.answer("Bu fandan darslar yo'q.", show_alert=True)
    await c.message.edit_reply_markup(reply_markup=markup)
    await c.answer()


@r.callback_query(F.data.startswith("vl:"))
async def vid_play(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    vid = int(c.data.split(":")[1])
    v = await db.one(
        "SELECT v.* FROM videos v JOIN group_videos gv ON gv.video_id=v.id WHERE gv.group_id=? AND v.id=?",
        (u["group_id"], vid),
    )
    if not v:
        return await c.answer("Bu dars sizga ochiq emas.", show_alert=True)
    await bot.send_video(c.from_user.id, v["file_id"], caption=esc(v["title"]), protect_content=True)
    await c.answer()


# ================================================================== O'QUVCHI: KITOBLAR
@r.message(F.text == B_BOOK)
async def book_menu(m: Message, state: FSMContext):
    await state.clear()
    if not await get_student(m):
        return
    await m.answer(
        "Fanni tanlang:",
        reply_markup=ikb([[("📐 Algebra", "bs:Algebra"), ("📏 Geometriya", "bs:Geometriya")], [("📚 Boshqa", "bs:Boshqa")]]),
    )


@r.callback_query(F.data.startswith("bs:"))
async def book_list(c: CallbackQuery):
    subj = c.data.split(":", 1)[1]
    rows = await db.fetch("SELECT id,title FROM books WHERE subject=? ORDER BY id", (subj,))
    if not rows:
        return await c.answer("Hozircha kitob yo'q.", show_alert=True)
    await c.message.answer(
        f"<b>{subj}</b> kitoblari:", reply_markup=ikb([[(row["title"][:60], f"bg:{row['id']}")] for row in rows])
    )
    await c.answer()


@r.callback_query(F.data.startswith("bg:"))
async def book_get(c: CallbackQuery):
    b = await db.one("SELECT * FROM books WHERE id=?", (int(c.data.split(":")[1]),))
    if b:
        await bot.send_document(c.from_user.id, b["file_id"], caption=esc(b["title"]))
    await c.answer()


# ================================================================== O'QUVCHI: TEST
@r.message(F.text == B_TEST)
async def test_menu(m: Message, state: FSMContext):
    await state.clear()
    if not await get_student(m):
        return
    await state.set_state(TestAccess.code)
    legacy = await db.count("SELECT COUNT(*) FROM tests WHERE access_code IS NULL OR access_code=''")
    markup = ikb([[("📋 Eski testlar ro'yxati", "tests:legacy")]]) if legacy else None
    await m.answer(
        "Testni ochish uchun admin bergan 4 xonali kodni kiriting:",
        reply_markup=markup,
    )


@r.callback_query(F.data == "tests:legacy")
async def legacy_tests(c: CallbackQuery, state: FSMContext):
    u = await get_student(c)
    if not u:
        return await c.answer()
    await state.clear()
    rows = await db.fetch(
        "SELECT id,title,mode FROM tests WHERE access_code IS NULL OR access_code='' ORDER BY id"
    )
    if not rows:
        return await c.answer("Kod bilan berilmagan testlar mavjud emas.", show_alert=True)
    keyboard = [
        [(f"{'🤖' if row['mode'] == 'auto' else '✍️'} {row['title']}"[:60], f"tp:{row['id']}")]
        for row in rows
    ]
    await c.message.answer("Eski testlardan birini tanlang:", reply_markup=ikb(keyboard))
    await c.answer()


@r.message(TestAccess.code, F.text)
async def test_access_code(m: Message, state: FSMContext):
    code = m.text.strip()
    if not re.fullmatch(r"\d{4}", code):
        return await m.answer("Kod 4 ta raqamdan iborat bo'lishi kerak. Qayta kiriting:")
    u = await get_student(m)
    if not u:
        return
    test = await db.one("SELECT * FROM tests WHERE access_code=?", (code,))
    if not test:
        return await m.answer("Bu kod bo'yicha test topilmadi. Kodni tekshirib qayta kiriting:")
    if test["mode"] == "auto" and not await db.count(
        "SELECT COUNT(*) FROM test_questions WHERE test_id=?", (test["id"],)
    ):
        return await m.answer("Bu avtomatik test hali savollar bilan to'ldirilmagan. Adminga murojaat qiling.")
    await state.clear()
    if test["mode"] == "auto":
        active = await db.one(
            "SELECT id FROM attempts WHERE test_id=? AND user_id=? AND finished_at IS NULL ORDER BY id DESC LIMIT 1",
            (test["id"], u["id"]),
        )
        action = ("▶️ Davom ettirish", f"quiz:resume:{active['id']}") if active else (
            "🚀 Testni ishlash", f"quiz:start:{test['id']}"
        )
        markup = ikb([[action]])
    else:
        last = await db.one(
            "SELECT status FROM submissions WHERE test_id=? AND user_id=? ORDER BY id DESC LIMIT 1",
            (test["id"], u["id"]),
        )
        status = {"approved": "✅ Tasdiqlangan", "pending": "⏳ Tekshirilmoqda",
                  "rejected": "❌ Xato, qayta topshiring"}.get(
            last["status"] if last else "", "🆕 Hali topshirilmagan"
        )
        markup = ikb([[("✍️ Javob yuborish", f"tsub:{test['id']}")]])
        await m.answer(f"Holat: {status}")
    if test["ctype"] in {"photo", "document"} and test["file_id"]:
        await send_content(
            m.chat.id, test["ctype"], test["file_id"], test["text"],
            caption=test["title"], markup=markup,
        )
    else:
        await m.answer(
            f"<b>{esc(test['title'])}</b>\n\n{esc(test['text'] or '')}",
            reply_markup=markup,
        )


@r.callback_query(F.data.startswith("tp:"))
async def test_open(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    tid = int(c.data.split(":")[1])
    t = await db.one("SELECT * FROM tests WHERE id=?", (tid,))
    if not t:
        return await c.answer("Test topilmadi.", show_alert=True)
    if t["mode"] == "auto":
        total = await db.count("SELECT COUNT(*) FROM test_questions WHERE test_id=?", (tid,))
        if not total:
            return await c.answer("Bu testda savollar yo'q.", show_alert=True)
        active = await db.one(
            "SELECT id FROM attempts WHERE test_id=? AND user_id=? AND finished_at IS NULL ORDER BY id DESC LIMIT 1",
            (tid, u["id"]),
        )
        action = ("▶️ Davom ettirish", f"quiz:resume:{active['id']}") if active else ("🚀 Testni boshlash", f"quiz:start:{tid}")
        await c.message.answer(
            f"🤖 <b>{esc(t['title'])}</b>\nSavollar soni: {total}\n"
            "Javoblarni tanlab, savollar orasida oldinga-orqaga o'tishingiz mumkin.",
            reply_markup=ikb([[action]]),
        )
        return await c.answer()
    last = await db.one(
        "SELECT status FROM submissions WHERE test_id=? AND user_id=? ORDER BY id DESC LIMIT 1", (tid, u["id"])
    )
    status = {"approved": "✅ Tasdiqlangan", "pending": "⏳ Tekshirilmoqda", "rejected": "❌ Xato, qayta topshiring"}.get(
        last["status"] if last else "", "🆕 Hali topshirilmagan"
    )
    await send_content(
        c.from_user.id, t["ctype"], t["file_id"], t["text"], caption=t["title"],
        markup=ikb([[("✍️ Javob yuborish", f"tsub:{tid}")]]),
    )
    await c.message.answer(f"Holat: {status}")
    await c.answer()


def quiz_keyboard(attempt_id, index, total, answers):
    selected = answers.get(str(index))
    choices = [[(f"{'✅ ' if selected == letter else ''}{letter}", f"qa:{attempt_id}:{index}:{letter}") for letter in "ABCD"]]
    nav = []
    if index > 0:
        nav.append(("⬅️ Oldingi", f"q:{attempt_id}:{index - 1}"))
    if index + 1 < total:
        nav.append(("Keyingi ➡️", f"q:{attempt_id}:{index + 1}"))
    rows = choices + ([nav] if nav else []) + [[("📋 Javoblarni ko'rish", f"qr:{attempt_id}")]]
    return ikb(rows)


QUIZ_SHEET_SIZE = 10


def quiz_sheet_keyboard(attempt_id, start, total, answers):
    rows = []
    end = min(start + QUIZ_SHEET_SIZE, total)
    for index in range(start, end):
        selected = answers.get(str(index))
        rows.append([
            (f"{'✅ ' if selected == letter else ''}{letter}", f"qa:{attempt_id}:{index}:{letter}")
            for letter in "ABCD"
        ])
    nav = []
    if start > 0:
        nav.append(("⬅️ Oldingi savollar", f"q:{attempt_id}:{start - QUIZ_SHEET_SIZE}"))
    if end < total:
        nav.append(("Keyingi savollar ➡️", f"q:{attempt_id}:{end}"))
    if nav:
        rows.append(nav)
    rows.append([("📋 Javoblarni tekshirish", f"qr:{attempt_id}")])
    return ikb(rows)


async def show_quiz_sheet(user_id, attempt_id, page, message, edit_existing=False):
    attempt = await db.one(
        "SELECT a.*,t.title FROM attempts a JOIN tests t ON t.id=a.test_id "
        "WHERE a.id=? AND a.user_id=? AND t.mode='auto' AND a.finished_at IS NULL",
        (attempt_id, user_id),
    )
    if not attempt:
        if message:
            await message.answer("Test urinishini topib bo'lmadi.")
        return
    total = await db.count(
        "SELECT COUNT(*) FROM test_questions WHERE test_id=?", (attempt["test_id"],)
    )
    start = page * QUIZ_SHEET_SIZE
    if page < 0 or start >= total:
        if message:
            await message.answer("Savollar sahifasi topilmadi.")
        return
    questions = await db.fetch(
        "SELECT position FROM test_questions WHERE test_id=? ORDER BY position,id LIMIT ? OFFSET ?",
        (attempt["test_id"], QUIZ_SHEET_SIZE, start),
    )
    answers = json.loads(attempt["answers"] or "{}")
    end = start + len(questions)
    lines = [
        f"{index + 1}. {esc(answers.get(str(index), '—'))}"
        for index in range(start, end)
    ]
    text = (
        f"📝 <b>{esc(attempt['title'])}</b>\n"
        f"Javoblarni belgilang ({start + 1}–{end} / {total}):\n\n"
        + "\n".join(lines)
    )
    markup = quiz_sheet_keyboard(attempt_id, start, total, answers)
    if edit_existing:
        await message.edit_text(text, reply_markup=markup)
    else:
        await message.answer(text, reply_markup=markup)


async def show_quiz_question(user_id, attempt_id, index, message=None, edit_existing=False):
    attempt = await db.one(
        "SELECT a.*,t.title FROM attempts a JOIN tests t ON t.id=a.test_id "
        "WHERE a.id=? AND a.user_id=? AND t.mode='auto'",
        (attempt_id, user_id),
    )
    if not attempt:
        if message:
            await message.answer("Test urinishini topib bo'lmadi.")
        return
    if attempt["finished_at"]:
        if message:
            await message.answer("Bu test yakunlangan.")
        return
    questions = await db.fetch(
        "SELECT id,position,file_id,text FROM test_questions WHERE test_id=? ORDER BY position,id",
        (attempt["test_id"],),
    )
    if index < 0 or index >= len(questions):
        if message:
            await message.answer("Savol topilmadi.")
        return
    question = questions[index]
    if not question["file_id"]:
        return await show_quiz_sheet(
            user_id, attempt_id, index // QUIZ_SHEET_SIZE, message, edit_existing
        )
    answers = json.loads(attempt["answers"] or "{}")
    caption = (
        f"<b>{esc(attempt['title'])}</b>\n"
        f"❓ {index + 1}/{len(questions)}-savol"
        + (f"\n\n{esc(question['text'])}" if question["text"] else "")
    )
    keyboard = quiz_keyboard(attempt_id, index, len(questions), answers)
    if not edit_existing:
        await message.answer_photo(question["file_id"], caption=caption, reply_markup=keyboard)
    else:
        await message.edit_media(
            media=InputMediaPhoto(media=question["file_id"], caption=caption),
            reply_markup=keyboard,
        )


@r.callback_query(F.data.startswith("quiz:start:"))
async def quiz_start(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    tid = int(c.data.split(":")[2])
    test = await db.one("SELECT id FROM tests WHERE id=? AND mode='auto'", (tid,))
    if not test:
        return await c.answer("Avtomatik test topilmadi.", show_alert=True)
    active = await db.one(
        "SELECT id FROM attempts WHERE test_id=? AND user_id=? AND finished_at IS NULL ORDER BY id DESC LIMIT 1",
        (tid, u["id"]),
    )
    attempt_id = active["id"] if active else await db.run(
        "INSERT INTO attempts(test_id,user_id,answers,started_at) VALUES(?,?,?,?)",
        (tid, u["id"], "{}", int(time.time())),
    )
    await c.answer()
    await show_quiz_question(c.from_user.id, attempt_id, 0, c.message)


@r.callback_query(F.data.startswith("quiz:resume:"))
async def quiz_resume(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    attempt_id = int(c.data.split(":")[2])
    attempt = await db.one(
        "SELECT id FROM attempts WHERE id=? AND user_id=? AND finished_at IS NULL",
        (attempt_id, u["id"]),
    )
    if not attempt:
        return await c.answer("Davom etadigan test topilmadi.", show_alert=True)
    await c.answer()
    await show_quiz_question(c.from_user.id, attempt_id, 0, c.message)


@r.callback_query(F.data.startswith("q:"))
async def quiz_navigate(c: CallbackQuery):
    _, attempt_id, index = c.data.split(":")
    await show_quiz_question(c.from_user.id, int(attempt_id), int(index), c.message, edit_existing=True)
    await c.answer()


@r.callback_query(F.data.startswith("qa:"))
async def quiz_answer(c: CallbackQuery):
    _, attempt_id, index, answer = c.data.split(":")
    attempt_id, index = int(attempt_id), int(index)
    if answer not in "ABCD" or len(answer) != 1:
        return await c.answer("Javob noto'g'ri.", show_alert=True)
    attempt = await db.one(
        "SELECT * FROM attempts WHERE id=? AND user_id=? AND finished_at IS NULL",
        (attempt_id, c.from_user.id),
    )
    if not attempt:
        return await c.answer("Test urinishini topib bo'lmadi.", show_alert=True)
    questions = await db.fetch(
        "SELECT id FROM test_questions WHERE test_id=? ORDER BY position,id", (attempt["test_id"],)
    )
    if index < 0 or index >= len(questions):
        return await c.answer("Savol topilmadi.", show_alert=True)
    answers = json.loads(attempt["answers"] or "{}")
    answers[str(index)] = answer
    await db.run("UPDATE attempts SET answers=? WHERE id=?", (json.dumps(answers), attempt_id))
    await show_quiz_question(
        c.from_user.id, attempt_id, index, c.message, edit_existing=True
    )
    await c.answer(f"{index + 1}-savol: {answer} tanlandi")


@r.callback_query(F.data.startswith("qr:"))
async def quiz_review(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    parts = c.data.split(":")
    attempt_id = int(parts[1])
    page = int(parts[2]) if len(parts) > 2 else 0
    attempt = await db.one(
        "SELECT a.*,t.title FROM attempts a JOIN tests t ON t.id=a.test_id "
        "WHERE a.id=? AND a.user_id=? AND a.finished_at IS NULL",
        (attempt_id, u["id"]),
    )
    if not attempt:
        return await c.answer("Test urinishini topib bo'lmadi.", show_alert=True)
    total = await db.count(
        "SELECT COUNT(*) FROM test_questions WHERE test_id=?", (attempt["test_id"],)
    )
    page_size = 20
    if page < 0 or page * page_size >= total:
        return await c.answer("Sahifa topilmadi.", show_alert=True)
    questions = await db.fetch(
        "SELECT id FROM test_questions WHERE test_id=? ORDER BY position,id LIMIT ? OFFSET ?",
        (attempt["test_id"], page_size, page * page_size),
    )
    answers = json.loads(attempt["answers"] or "{}")
    start = page * page_size
    lines = [f"{i + 1}. {esc(answers.get(str(i), '—'))}" for i in range(start, start + len(questions))]
    buttons = [
        (f"{i + 1}{' ✅ ' + answers[str(i)] if str(i) in answers else ' ❔'}", f"qe:{attempt_id}:{i}")
        for i in range(start, start + len(questions))
    ]
    rows = [buttons[i:i + 4] for i in range(0, len(buttons), 4)]
    if page > 0:
        rows.append([("⬅️ Oldingi savollar", f"qr:{attempt_id}:{page - 1}")])
    if start + page_size < total:
        rows.append([("Keyingi savollar ➡️", f"qr:{attempt_id}:{page + 1}")])
    unanswered = total - len(answers)
    if not unanswered:
        rows.append([("✅ Testni tugatish", f"qf:{attempt_id}")])
    text = (
        f"📋 <b>{esc(attempt['title'])} — javoblaringiz</b>\n"
        f"Savollar {start + 1}–{min(start + len(questions), total)} / {total}\n\n"
        + "\n".join(lines)
    )
    if unanswered:
        text += f"\n\n❗ {unanswered} ta savolga javob berilmagan. O'zgartirish uchun savol raqamini bosing."
    await c.message.answer(text, reply_markup=ikb(rows))
    await c.answer()


@r.callback_query(F.data.startswith("qe:"))
async def quiz_edit_answer(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    _, attempt_id, index = c.data.split(":")
    await c.answer()
    await show_quiz_question(c.from_user.id, int(attempt_id), int(index), c.message)


@r.callback_query(F.data.startswith("qf:"))
async def quiz_finish(c: CallbackQuery):
    u = await get_student(c)
    if not u:
        return await c.answer()
    attempt_id = int(c.data.split(":")[1])
    attempt = await db.one(
        "SELECT a.*,t.title FROM attempts a JOIN tests t ON t.id=a.test_id "
        "WHERE a.id=? AND a.user_id=? AND t.mode='auto'",
        (attempt_id, u["id"]),
    )
    if not attempt or attempt["finished_at"]:
        return await c.answer("Bu test allaqachon yakunlangan yoki topilmadi.", show_alert=True)
    questions = await db.fetch(
        "SELECT position,correct FROM test_questions WHERE test_id=? ORDER BY position,id",
        (attempt["test_id"],),
    )
    answers = json.loads(attempt["answers"] or "{}")
    if len(answers) != len(questions) or any(str(i) not in answers for i in range(len(questions))):
        return await c.answer("Avval barcha savollarga javob bering.", show_alert=True)
    score = sum(
        1 for i, question in enumerate(questions)
        if answers[str(i)].upper() == (question["correct"] or "").upper()
    )
    finished_at = int(time.time())
    await db.run(
        "UPDATE attempts SET score=?,total=?,finished_at=? WHERE id=? AND finished_at IS NULL",
        (score, len(questions), finished_at, attempt_id),
    )
    result = f"📊 «{esc(attempt['title'])}» testi natijasi: <b>{score}/{len(questions)}</b>"
    await c.message.answer(f"✅ Test yakunlandi!\n{result}")
    await c.answer()
    admin_text = (
        f"📊 <b>Avtomatik test yakunlandi</b>\n"
        f"👤 {esc(u['full_name'])}\n"
        f"📝 {esc(attempt['title'])}\n"
        f"Natija: <b>{score}/{len(questions)}</b>\n\n"
        + "\n".join(f"{i + 1}. {esc(answers[str(i)])}" for i in range(len(questions)))
    )
    for admin in await db.fetch("SELECT id FROM admins"):
        await safe_send(admin["id"], admin_text)


@r.callback_query(F.data.startswith("tsub:"))
async def test_submit_start(c: CallbackQuery, state: FSMContext):
    u = await get_student(c)
    if not u:
        return await c.answer()
    tid = int(c.data.split(":")[1])
    test = await db.one("SELECT mode FROM tests WHERE id=?", (tid,))
    if not test or test["mode"] != "manual":
        return await c.answer("Bu avtomatik test. Uni testlar ro'yxatidan boshlang.", show_alert=True)
    last = await db.one(
        "SELECT status FROM submissions WHERE test_id=? AND user_id=? ORDER BY id DESC LIMIT 1", (tid, u["id"])
    )
    if last and last["status"] == "approved":
        return await c.answer("Bu test allaqachon tasdiqlangan ✅", show_alert=True)
    if last and last["status"] == "pending":
        return await c.answer("Javobingiz tekshirilmoqda ⏳", show_alert=True)
    await state.set_state(Submit.photos)
    await state.update_data(tid=tid, photos=[])
    await c.message.answer(
        "📸 Yechilgan testni rasmga olib yuboring (bir nechta rasm mumkin, 10 tagacha).",
        reply_markup=ikb([[("❌ Bekor qilish", "tcancel")]]),
    )
    await c.answer()


@r.message(Submit.photos, F.photo)
async def submit_photo(m: Message, state: FSMContext):
    d = await state.get_data()
    photos = d.get("photos", [])
    if len(photos) >= 10:
        return await m.answer("Eng ko'pi bilan 10 ta rasm. «Yuborish» ni bosing.")
    photos.append(m.photo[-1].file_id)
    await state.update_data(photos=photos)
    await m.answer(
        f"📎 {len(photos)} ta rasm qabul qilindi.",
        reply_markup=ikb([[(f"✅ Yuborish ({len(photos)})", "tdone")], [("❌ Bekor qilish", "tcancel")]]),
    )


@r.message(Submit.photos)
async def submit_other(m: Message):
    await m.answer("Iltimos, faqat rasm yuboring.")


@r.callback_query(F.data == "tcancel")
async def submit_cancel(c: CallbackQuery, state: FSMContext):
    await state.clear()
    await c.message.answer("Bekor qilindi.")
    await c.answer()


@r.callback_query(F.data == "tdone", Submit.photos)
async def submit_done(c: CallbackQuery, state: FSMContext):
    d = await state.get_data()
    photos = d.get("photos", [])
    if not photos:
        return await c.answer("Avval rasm yuboring.", show_alert=True)
    u = await db.one("SELECT * FROM users WHERE id=?", (c.from_user.id,))
    t = await db.one("SELECT * FROM tests WHERE id=?", (d["tid"],))
    g = await db.one("SELECT name FROM groups WHERE id=?", (u["group_id"],))
    sid = await db.run(
        "INSERT INTO submissions(test_id,user_id,status,created_at) VALUES(?,?,?,?)",
        (d["tid"], u["id"], "pending", int(time.time())),
    )
    for f in photos:
        await db.run("INSERT INTO sub_files(sub_id,file_id) VALUES(?,?)", (sid, f))
    await state.clear()
    await c.message.answer("✅ Javobingiz adminlarga yuborildi. Natijani kuting.")
    await c.answer()

    info = (
        f"📥 <b>Yangi test topshirildi</b>\n👤 {esc(u['full_name'])}\n👥 {esc(g['name'])}\n"
        f"📝 {esc(t['title'])}\n📞 {esc(u['phone'] or '-')}"
    )
    for a in await db.fetch("SELECT id FROM admins"):
        try:
            if len(photos) == 1:
                await bot.send_photo(a["id"], photos[0], caption=info)
            else:
                mg = MediaGroupBuilder(caption=info)
                for f in photos:
                    mg.add_photo(f)
                await bot.send_media_group(a["id"], mg.build())
            await bot.send_message(
                a["id"], "Natijani belgilang:",
                reply_markup=ikb([[("✅ Tasdiqlash", f"sa:{sid}"), ("❌ Xato", f"sr:{sid}")]]),
            )
        except Exception as e:
            logging.warning("admin notify fail %s: %s", a["id"], e)


# ---- admin tekshiruvi
@r.callback_query(F.data.startswith("sa:"), IsAdmin())
async def sub_approve(c: CallbackQuery):
    sid = int(c.data.split(":")[1])
    s = await db.one("SELECT s.*, t.title FROM submissions s JOIN tests t ON t.id=s.test_id WHERE s.id=?", (sid,))
    if not s or s["status"] != "pending":
        return await c.answer("Bu javob allaqachon ko'rib chiqilgan.", show_alert=True)
    await db.run("UPDATE submissions SET status='approved', reviewed_by=? WHERE id=?", (c.from_user.id, sid))
    await c.message.edit_text("✅ Tasdiqlandi")
    await safe_send(s["user_id"], f"✅ «{esc(s['title'])}» testingiz tasdiqlandi. Barakalla! Keyingi darsga o'tishingiz mumkin.")
    await c.answer()


@r.callback_query(F.data.startswith("sr:"), IsAdmin())
async def sub_reject(c: CallbackQuery, state: FSMContext):
    sid = int(c.data.split(":")[1])
    s = await db.one("SELECT status FROM submissions WHERE id=?", (sid,))
    if not s or s["status"] != "pending":
        return await c.answer("Bu javob allaqachon ko'rib chiqilgan.", show_alert=True)
    await state.set_state(Review.comment)
    await state.update_data(sid=sid)
    await c.message.answer("✍️ Xato izohini yozing (nimalar noto'g'ri):")
    await c.answer()


@r.message(Review.comment, F.text)
async def sub_reject_comment(m: Message, state: FSMContext):
    d = await state.get_data()
    sid = d["sid"]
    s = await db.one("SELECT s.*, t.title FROM submissions s JOIN tests t ON t.id=s.test_id WHERE s.id=?", (sid,))
    await state.clear()
    if not s or s["status"] != "pending":
        return await m.answer("Bu javob allaqachon ko'rib chiqilgan.")
    await db.run(
        "UPDATE submissions SET status='rejected', comment=?, reviewed_by=? WHERE id=?",
        (m.text, m.from_user.id, sid),
    )
    await safe_send(
        s["user_id"],
        f"❌ «{esc(s['title'])}» testida xato bor:\n\n{esc(m.text)}\n\nQayta ishlab topshiring.",
        ikb([[("✍️ Qayta topshirish", f"tsub:{s['test_id']}")]]),
    )
    await m.answer("Xato izohi o'quvchiga yuborildi ✅")


# ================================================================== ADMIN: GURUHLAR
async def groups_text_kb():
    rows = await db.fetch(
        "SELECT g.id, g.name, (SELECT COUNT(*) FROM users u WHERE u.group_id=g.id) AS n FROM groups g ORDER BY g.id"
    )
    kb = [[(f"👁 {row['name']} ({row['n']})", f"g:view:{row['id']}")] for row in rows]
    kb.append([("➕ Yangi guruh", "g:new")])
    return ("👥 <b>Guruhlar</b>" if rows else "Hozircha guruhlar yo'q."), ikb(kb)


@r.message(F.text == A_GROUPS, IsAdmin())
async def groups_list(m: Message, state: FSMContext):
    await state.clear()
    text, kb = await groups_text_kb()
    await m.answer(text, reply_markup=kb)


@r.callback_query(F.data == "g:new", IsAdmin())
async def group_new(c: CallbackQuery, state: FSMContext):
    await state.set_state(GrpNew.name)
    await c.message.answer("Yangi guruh nomini yozing:")
    await c.answer()


@r.message(GrpNew.name, F.text)
async def group_new_name(m: Message, state: FSMContext):
    name = m.text.strip()
    if not name:
        return await m.answer("Guruh nomi bo'sh bo'lmasligi kerak.")
    try:
        await db.run("INSERT INTO groups(name) VALUES(?)", (name,))
        await m.answer(f"✅ «{esc(name)}» guruhi yaratildi.")
    except sqlite3.IntegrityError:
        await m.answer("Bunday nomli guruh bor.")
    else:
        await state.clear()


async def group_view_render(gid):
    g = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    if not g:
        return "Guruh topilmadi.", None
    studs = await db.fetch("SELECT id, full_name FROM users WHERE group_id=? ORDER BY full_name", (gid,))
    text = f"👥 <b>{esc(g['name'])}</b> — {len(studs)} ta o'quvchi\n(🗑 bosilsa guruhdan chiqariladi)"
    kb = [[(f"🗑 {s['full_name']}"[:60], f"g:rm:{gid}:{s['id']}")] for s in studs]
    kb.extend([
        [("✏️ Guruh nomini tahrirlash", f"g:edit:{gid}")],
        [("🗑 Guruhni o'chirish", f"g:del:ask:{gid}")],
    ])
    return text, ikb(kb)


@r.callback_query(F.data.startswith("g:view:"), IsAdmin())
async def group_view(c: CallbackQuery):
    gid = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM groups WHERE id=?", (gid,)):
        return await c.answer("Guruh topilmadi.", show_alert=True)
    text, kb = await group_view_render(gid)
    await c.message.answer(text, reply_markup=kb)
    await c.answer()


@r.callback_query(F.data.startswith("g:rm:"), IsAdmin())
async def group_remove(c: CallbackQuery):
    _, _, gid, uid = c.data.split(":")
    if not await db.one("SELECT 1 FROM users WHERE id=? AND group_id=?", (int(uid), int(gid))):
        return await c.answer("O'quvchi bu guruhda topilmadi.", show_alert=True)
    await db.run("UPDATE users SET group_id=NULL, pending_group=NULL WHERE id=?", (int(uid),))
    text, kb = await group_view_render(int(gid))
    await c.message.edit_text(text, reply_markup=kb)
    await c.answer("Chiqarildi")


@r.callback_query(F.data.startswith("g:edit:"), IsAdmin())
async def group_edit_start(c: CallbackQuery, state: FSMContext):
    gid = int(c.data.split(":")[2])
    group = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    if not group:
        return await c.answer("Guruh topilmadi.", show_alert=True)
    await state.set_state(GrpEdit.name)
    await state.update_data(group_id=gid)
    await c.message.answer(f"«{esc(group['name'])}» guruhining yangi nomini yozing:")
    await c.answer()


@r.message(GrpEdit.name, F.text, IsAdmin())
async def group_edit_name(m: Message, state: FSMContext):
    name = m.text.strip()
    if not name:
        return await m.answer("Guruh nomi bo'sh bo'lmasligi kerak.")
    data = await state.get_data()
    try:
        await db.run("UPDATE groups SET name=? WHERE id=?", (name, data["group_id"]))
    except sqlite3.IntegrityError:
        return await m.answer("Bunday nomli guruh bor. Boshqa nom kiriting.")
    await state.clear()
    await m.answer(f"✅ Guruh nomi «{esc(name)}» qilib o'zgartirildi.")


@r.callback_query(F.data.startswith("g:del:ask:"), IsAdmin())
async def group_delete_confirm(c: CallbackQuery):
    gid = int(c.data.split(":")[3])
    group = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    if not group:
        return await c.answer("Guruh topilmadi.", show_alert=True)
    await c.message.answer(
        f"«{esc(group['name'])}» guruhini o'chirasizmi? O'quvchilar o'chmaydi, faqat guruhdan chiqariladi.",
        reply_markup=ikb([
            [("🗑 Ha, guruhni o'chirish", f"g:del:yes:{gid}")],
            [("Bekor qilish", "g:del:no")],
        ]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("g:del:yes:"), IsAdmin())
async def group_delete(c: CallbackQuery):
    gid = int(c.data.split(":")[3])
    if not await db.one("SELECT 1 FROM groups WHERE id=?", (gid,)):
        return await c.answer("Guruh topilmadi.", show_alert=True)
    await db.run("UPDATE users SET group_id=NULL, pending_group=NULL WHERE group_id=? OR pending_group=?", (gid, gid))
    await db.run("DELETE FROM group_videos WHERE group_id=?", (gid,))
    await db.run("DELETE FROM groups WHERE id=?", (gid,))
    await c.message.edit_text("✅ Guruh o'chirildi. O'quvchilar guruhsiz qoldi.")
    await c.answer()


@r.callback_query(F.data == "g:del:no", IsAdmin())
async def group_delete_cancel(c: CallbackQuery):
    await c.message.edit_text("Guruhni o'chirish bekor qilindi.")
    await c.answer()


# ================================================================== ADMIN: O'QUVCHILAR RO'YXATI
STUDENTS_PAGE_SIZE = 10


async def students_page(offset):
    total = await db.count("SELECT COUNT(*) FROM users")
    rows = await db.fetch(
        "SELECT u.id, u.full_name, u.phone, u.username, u.blocked, u.created_at, u.last_active, "
        "g.name AS group_name FROM users u LEFT JOIN groups g ON g.id=u.group_id "
        "ORDER BY u.created_at DESC, u.id DESC LIMIT ? OFFSET ?",
        (STUDENTS_PAGE_SIZE, offset),
    )
    if not rows:
        return "Hozircha o'quvchilar yo'q.", None

    pages = (total + STUDENTS_PAGE_SIZE - 1) // STUDENTS_PAGE_SIZE
    page = offset // STUDENTS_PAGE_SIZE + 1
    entries = []
    for row in rows:
        full_name = str(row["full_name"] or "Ism ko'rsatilmagan")[:80]
        created = time.strftime("%d.%m.%Y", time.localtime(row["created_at"])) if row["created_at"] else "—"
        active = time.strftime("%d.%m.%Y %H:%M", time.localtime(row["last_active"])) if row["last_active"] else "—"
        username = f"@{str(row['username'])[:40]}" if row["username"] else "—"
        group_name = str(row["group_name"] or "Guruhga biriktirilmagan")[:50]
        status = "Botni bloklagan" if row["blocked"] else "Faol"
        entries.append(
            f"👤 <b>{esc(full_name)}</b>\n"
            f"🆔 <code>{row['id']}</code> | 📞 {esc(str(row['phone'] or '—')[:32])}\n"
            f"🔗 {esc(username)} | 👥 {esc(group_name)}\n"
            f"📌 {status} | Ro'yxatdan o'tgan: {created}\n"
            f"🕒 Oxirgi faollik: {active}"
        )

    navigation = []
    controls = [
        [(f"✏️ {str(row['full_name'] or 'Oquvchi')[:40]}", f"student:edit:{row['id']}")]
        for row in rows
    ]
    if offset > 0:
        navigation.append(("⬅️ Oldingi", f"students:p:{offset - STUDENTS_PAGE_SIZE}"))
    if offset + STUDENTS_PAGE_SIZE < total:
        navigation.append(("Keyingi ➡️", f"students:p:{offset + STUDENTS_PAGE_SIZE}"))
    keyboard = controls
    if navigation:
        keyboard.append(navigation)
    markup = ikb(keyboard)
    return f"👨‍🎓 <b>O'quvchilar</b> — {total} ta (sahifa {page}/{pages})\n\n" + "\n\n".join(entries), markup


@r.message(F.text == A_STUDENTS, IsAdmin())
async def students_list(m: Message, state: FSMContext):
    await state.clear()
    text, markup = await students_page(0)
    await m.answer(text, reply_markup=markup)


@r.callback_query(F.data.startswith("students:p:"), IsAdmin())
async def students_page_callback(c: CallbackQuery):
    offset = int(c.data.split(":")[2])
    if offset < 0:
        return await c.answer("Sahifa topilmadi.", show_alert=True)
    text, markup = await students_page(offset)
    await c.message.edit_text(text, reply_markup=markup)
    await c.answer()


@r.callback_query(F.data.startswith("student:edit:"), IsAdmin())
async def student_edit_menu(c: CallbackQuery):
    uid = int(c.data.split(":")[2])
    student = await db.one("SELECT full_name FROM users WHERE id=?", (uid,))
    if not student:
        return await c.answer("O'quvchi topilmadi.", show_alert=True)
    await c.message.answer(
        f"<b>{esc(student['full_name'] or 'Oquvchi')}</b> ma'lumotini o'zgartiring:",
        reply_markup=ikb([
            [("✏️ Ism-familiya", f"student:name:{uid}"), ("📞 Telefon", f"student:phone:{uid}")],
            [("👥 Guruh", f"student:group:{uid}")],
        ]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("student:name:"), IsAdmin())
async def student_edit_name_start(c: CallbackQuery, state: FSMContext):
    uid = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM users WHERE id=?", (uid,)):
        return await c.answer("O'quvchi topilmadi.", show_alert=True)
    await state.set_state(StudentEdit.name)
    await state.update_data(user_id=uid)
    await c.message.answer("O'quvchining yangi ism-familiyasini yozing:")
    await c.answer()


@r.message(StudentEdit.name, F.text, IsAdmin())
async def student_edit_name(m: Message, state: FSMContext):
    name = m.text.strip()
    if len(name) < 3:
        return await m.answer("Ism-familiyani kamida 3 ta belgi bilan yozing.")
    data = await state.get_data()
    await db.run("UPDATE users SET full_name=? WHERE id=?", (name, data["user_id"]))
    await state.clear()
    await m.answer("✅ O'quvchi ismi o'zgartirildi.")


@r.callback_query(F.data.startswith("student:phone:"), IsAdmin())
async def student_edit_phone_start(c: CallbackQuery, state: FSMContext):
    uid = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM users WHERE id=?", (uid,)):
        return await c.answer("O'quvchi topilmadi.", show_alert=True)
    await state.set_state(StudentEdit.phone)
    await state.update_data(user_id=uid)
    await c.message.answer("Yangi telefon raqamini yozing:")
    await c.answer()


@r.message(StudentEdit.phone, F.text, IsAdmin())
async def student_edit_phone(m: Message, state: FSMContext):
    phone = m.text.strip()
    if not phone or len(phone) > 32:
        return await m.answer("Telefon raqami bo'sh bo'lmasin va 32 belgidan oshmasin.")
    data = await state.get_data()
    await db.run("UPDATE users SET phone=? WHERE id=?", (phone, data["user_id"]))
    await state.clear()
    await m.answer("✅ Telefon raqami o'zgartirildi.")


@r.callback_query(F.data.startswith("student:group:"), IsAdmin())
async def student_edit_group_menu(c: CallbackQuery):
    uid = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM users WHERE id=?", (uid,)):
        return await c.answer("O'quvchi topilmadi.", show_alert=True)
    groups = await db.fetch("SELECT id, name FROM groups ORDER BY name")
    keyboard = [[(row["name"][:50], f"student:g:{uid}:{row['id']}")] for row in groups]
    keyboard.append([("Guruhdan chiqarish", f"student:ungroup:{uid}")])
    await c.message.answer(
        "O'quvchini qaysi guruhga o'tkazamiz?",
        reply_markup=ikb(keyboard),
    )
    await c.answer()


@r.callback_query(F.data.startswith("student:g:"), IsAdmin())
async def student_edit_group(c: CallbackQuery):
    _, _, uid, gid = c.data.split(":")
    uid, gid = int(uid), int(gid)
    student = await db.one("SELECT full_name FROM users WHERE id=?", (uid,))
    group = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    if not student or not group:
        return await c.answer("O'quvchi yoki guruh topilmadi.", show_alert=True)
    await db.run("UPDATE users SET pending_group=? WHERE id=?", (gid, uid))
    ok = await safe_send(
        uid,
        f"Admin sizni <b>{esc(group['name'])}</b> guruhiga o'tkazmoqchi. Rozimisiz?",
        ikb([[("Roziman ✓", f"agree:{gid}"), ("Rad etish", "decline")]]),
    )
    await c.answer("So'rov yuborildi." if ok else "O'quvchiga xabar yuborib bo'lmadi.", show_alert=not ok)


@r.callback_query(F.data.startswith("student:ungroup:"), IsAdmin())
async def student_ungroup(c: CallbackQuery):
    uid = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM users WHERE id=?", (uid,)):
        return await c.answer("O'quvchi topilmadi.", show_alert=True)
    await db.run("UPDATE users SET group_id=NULL, pending_group=NULL WHERE id=?", (uid,))
    await c.answer("O'quvchi guruhdan chiqarildi.")


# ================================================================== ADMIN: O'QUVCHI QO'SHISH
@r.message(F.text == A_ADDST, IsAdmin())
async def add_student(m: Message, state: FSMContext):
    await state.clear()
    rows = await db.fetch("SELECT id,name FROM groups ORDER BY id")
    if not rows:
        return await m.answer("Avval «👥 Guruhlar» bo'limida guruh yarating.")
    await m.answer("Qaysi guruhga qo'shamiz?", reply_markup=ikb([[(row["name"], f"as:g:{row['id']}")] for row in rows]))


@r.callback_query(F.data.startswith("as:g:"), IsAdmin())
async def add_student_pick(c: CallbackQuery):
    gid = int(c.data.split(":")[2])
    studs = await db.fetch(
        "SELECT id, full_name FROM users WHERE group_id IS NULL AND blocked=0 ORDER BY created_at DESC LIMIT 60"
    )
    if not studs:
        return await c.answer("Guruhsiz o'quvchi yo'q. O'quvchi avval botda /start bosishi kerak.", show_alert=True)
    await c.message.answer(
        "O'quvchini tanlang:",
        reply_markup=ikb([[(s["full_name"][:50], f"as:u:{gid}:{s['id']}")] for s in studs]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("as:u:"), IsAdmin())
async def add_student_send(c: CallbackQuery):
    _, _, gid, uid = c.data.split(":")
    gid, uid = int(gid), int(uid)
    g = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    if not g or not await db.one("SELECT 1 FROM users WHERE id=?", (uid,)):
        return await c.answer("O'quvchi yoki guruh topilmadi.", show_alert=True)
    await db.run("UPDATE users SET pending_group=? WHERE id=?", (gid, uid))
    ok = await safe_send(
        uid,
        f"Admin sizni <b>{esc(g['name'])}</b> guruhiga qo'shmoqchi. Rozimisiz?",
        ikb([[("Roziman ✓", f"agree:{gid}"), ("Rad etish", "decline")]]),
    )
    await c.message.answer("📨 So'rov o'quvchiga yuborildi." if ok else "❌ O'quvchiga xabar yuborib bo'lmadi.")
    await c.answer()


@r.callback_query(F.data.startswith("agree:"))
async def student_agree(c: CallbackQuery):
    gid = int(c.data.split(":")[1])
    u = await db.one("SELECT pending_group FROM users WHERE id=?", (c.from_user.id,))
    if not u or u["pending_group"] != gid:
        return await c.answer("So'rov eskirgan.", show_alert=True)
    g = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    await db.run("UPDATE users SET group_id=?, pending_group=NULL WHERE id=?", (gid, c.from_user.id))
    await c.message.edit_text(f"✅ Siz <b>{esc(g['name'])}</b> guruhiga qo'shildingiz!")
    await c.message.answer("Darslarni boshlashingiz mumkin 👇", reply_markup=await menu_for(c.from_user.id))
    await c.answer()


@r.callback_query(F.data == "decline")
async def student_decline(c: CallbackQuery):
    await db.run("UPDATE users SET pending_group=NULL WHERE id=?", (c.from_user.id,))
    await c.message.edit_text("Rad etildi.")
    await c.answer()


# ================================================================== ADMIN: VIDEO KUTUBXONA
async def save_video(m: Message):
    cap = (m.caption or "").strip()
    subject, title = "Boshqa", cap or "Nomsiz video"
    if "|" in cap:
        a, b = [x.strip() for x in cap.split("|", 1)]
        for s in SUBJECTS:
            if a.lower().startswith(s.lower()[:3]):
                subject, title = s, b or s
                break
    await db.run(
        "INSERT INTO videos(file_id,title,subject,created_at) VALUES(?,?,?,?)",
        (m.video.file_id, title, subject, int(time.time())),
    )
    await m.reply(f"✅ Videolar bazasiga saqlandi: <b>{esc(subject)}</b> — {esc(title)}")


@r.message(F.video, F.chat.id == VIDEO_GROUP_ID)
async def video_from_group(m: Message):
    await save_video(m)


@r.message(F.video, F.chat.type == "private", IsAdmin())
async def video_from_admin(m: Message, state: FSMContext):
    if await state.get_state() == VideoAdd.upload.state:
        title = (m.caption or "").strip() or "Nomsiz video"
        await state.update_data(video_file_id=m.video.file_id, video_title=title)
        await m.answer(
            f"«{esc(title)}» videosi qaysi fanga qo'shilsin?",
            reply_markup=ikb([
                [("📐 Algebra", "video:add:Algebra"), ("📏 Geometriya", "video:add:Geometriya")],
                [("📐📏 Algebra va Geometriya", "video:add:both")],
            ]),
        )
    elif await state.get_state() is None:
        await save_video(m)
    else:
        await m.answer("Video qo'shish bosqichini «Video qo'shish» tugmasi orqali boshlang.")


@r.message(F.text == A_VADD, IsAdmin())
async def video_add_start(m: Message, state: FSMContext):
    await state.clear()
    await state.set_state(VideoAdd.upload)
    await m.answer(
        "Video darsni izohi (caption) bilan yuboring. Izoh video nomi sifatida saqlanadi."
    )


@r.message(VideoAdd.upload, IsAdmin())
async def video_add_invalid(m: Message):
    await m.answer("Bu bo'limda video fayl yuboring. Izoh video nomi bo'lib saqlanadi.")


@r.callback_query(F.data.startswith("video:add:"), IsAdmin(), VideoAdd.upload)
async def video_add_subject(c: CallbackQuery, state: FSMContext):
    subject = c.data.split(":")[2]
    subjects = ["Algebra", "Geometriya"] if subject == "both" else [subject]
    if any(item not in SUBJECTS for item in subjects):
        return await c.answer("Fan tanlovi noto'g'ri.", show_alert=True)
    data = await state.get_data()
    if not data.get("video_file_id"):
        return await c.answer("Avval video yuboring.", show_alert=True)
    for selected_subject in subjects:
        await db.run(
            "INSERT INTO videos(file_id,title,subject,created_at) VALUES(?,?,?,?)",
            (data["video_file_id"], data["video_title"], selected_subject, int(time.time())),
        )
    await state.clear()
    await c.message.answer(
        f"✅ «{esc(data['video_title'])}» video darsi "
        f"{' va '.join(subjects)} faniga saqlandi. Yana video qo'shish uchun «{A_VADD}» ni bosing."
    )
    await c.answer()


VIDEO_PAGE_SIZE = 8


async def picker_kb(gid, page=0):
    total = await db.count("SELECT COUNT(*) FROM videos")
    vids = await db.fetch(
        "SELECT id,title,subject FROM videos ORDER BY id DESC LIMIT ? OFFSET ?",
        (VIDEO_PAGE_SIZE, page * VIDEO_PAGE_SIZE),
    )
    att = {x["video_id"] for x in await db.fetch("SELECT video_id FROM group_videos WHERE group_id=?", (gid,))}
    kb = []
    for v in vids:
        mark = "✅ " if v["id"] in att else "▫️ "
        kb.append([((mark + f"{v['subject']}: {v['title']}")[:58], f"va:t:{gid}:{v['id']}")])
    nav = []
    if page > 0:
        nav.append(("⬅️ Oldingi", f"va:p:{gid}:{page - 1}"))
    if (page + 1) * VIDEO_PAGE_SIZE < total:
        nav.append(("Keyingi ➡️", f"va:p:{gid}:{page + 1}"))
    if nav:
        kb.append(nav)
    kb.append([("☑️ Hammasi", f"va:all:{gid}:{page}"), ("🏁 Tayyor", "va:done")])
    return ikb(kb)


async def video_delete_kb(page=0):
    total = await db.count("SELECT COUNT(*) FROM videos")
    videos = await db.fetch(
        "SELECT id,title,subject FROM videos ORDER BY id DESC LIMIT ? OFFSET ?",
        (VIDEO_PAGE_SIZE, page * VIDEO_PAGE_SIZE),
    )
    keyboard = [
        [(
            f"{video['subject']}: {video['title']}"[:58],
            f"vdel:ask:{video['id']}:{page}",
        )]
        for video in videos
    ]
    nav = []
    if page > 0:
        nav.append(("⬅️ Oldingi", f"vdel:p:{page - 1}"))
    if (page + 1) * VIDEO_PAGE_SIZE < total:
        nav.append(("Keyingi ➡️", f"vdel:p:{page + 1}"))
    if nav:
        keyboard.append(nav)
    return ikb(keyboard)


@r.message(F.text == A_VDEL, IsAdmin())
async def video_delete_menu(m: Message, state: FSMContext):
    await state.clear()
    if not await db.count("SELECT COUNT(*) FROM videos"):
        return await m.answer("O'chiriladigan video mavjud emas.")
    await m.answer("O'chirmoqchi bo'lgan videoni tanlang:", reply_markup=await video_delete_kb())


@r.callback_query(F.data.startswith("vdel:p:"), IsAdmin())
async def video_delete_page(c: CallbackQuery):
    page = int(c.data.split(":")[2])
    if page < 0:
        return await c.answer("Sahifa topilmadi.", show_alert=True)
    await c.message.edit_reply_markup(reply_markup=await video_delete_kb(page))
    await c.answer()


@r.callback_query(F.data.startswith("vdel:ask:"), IsAdmin())
async def video_delete_confirm(c: CallbackQuery):
    _, _, video_id, page = c.data.split(":")
    video = await db.one("SELECT title,subject FROM videos WHERE id=?", (int(video_id),))
    if not video:
        return await c.answer("Video topilmadi.", show_alert=True)
    await c.message.answer(
        f"«{esc(video['title'])}» ({esc(video['subject'])}) videosini o'chirasizmi?",
        reply_markup=ikb([
            [("🗑 Ha, o'chirish", f"vdel:yes:{video_id}:{page}")],
            [("Bekor qilish", f"vdel:cancel:{page}")],
        ]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("vdel:yes:"), IsAdmin())
async def video_delete(c: CallbackQuery):
    _, _, video_id, page = c.data.split(":")
    video_id, page = int(video_id), int(page)
    if not await db.one("SELECT 1 FROM videos WHERE id=?", (video_id,)):
        return await c.answer("Video topilmadi.", show_alert=True)
    await db.run("DELETE FROM group_videos WHERE video_id=?", (video_id,))
    await db.run("DELETE FROM videos WHERE id=?", (video_id,))
    await c.message.edit_text("✅ Video va guruhlarga biriktirilgan havolalari o'chirildi.")
    total = await db.count("SELECT COUNT(*) FROM videos")
    page = min(page, max(0, (total - 1) // VIDEO_PAGE_SIZE))
    await c.message.answer("Video kutubxonasi:", reply_markup=await video_delete_kb(page))
    await c.answer()


@r.callback_query(F.data.startswith("vdel:cancel:"), IsAdmin())
async def video_delete_cancel(c: CallbackQuery):
    page = int(c.data.split(":")[2])
    await c.message.edit_text("Video o'chirish bekor qilindi.")
    await c.message.answer("Video kutubxonasi:", reply_markup=await video_delete_kb(page))
    await c.answer()


@r.message(F.text == A_VATT, IsAdmin())
async def vid_attach(m: Message, state: FSMContext):
    await state.clear()
    if not await db.count("SELECT COUNT(*) FROM videos"):
        return await m.answer(
            "Videolar bazasi bo'sh. «Video qo'shish» tugmasini bosib, videoni izohi bilan yuboring."
        )
    rows = await db.fetch("SELECT id,name FROM groups ORDER BY id")
    if not rows:
        return await m.answer("Avval guruh yarating.")
    await m.answer("Qaysi guruhga video biriktiramiz?", reply_markup=ikb([[(r_["name"], f"va:g:{r_['id']}")] for r_ in rows]))


@r.callback_query(F.data.startswith("va:g:"), IsAdmin())
async def vid_attach_group(c: CallbackQuery):
    gid = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM groups WHERE id=?", (gid,)):
        return await c.answer("Guruh topilmadi.", show_alert=True)
    await c.message.answer(
        "Videoni bosing: ✅ — biriktirilgan, ▫️ — biriktirilmagan. "
        "Sahifalar orqali barcha videolarni ko'rishingiz mumkin.",
        reply_markup=await picker_kb(gid),
    )
    await c.answer()


@r.callback_query(F.data.startswith("va:p:"), IsAdmin())
async def vid_attach_page(c: CallbackQuery):
    _, _, gid, page = c.data.split(":")
    gid, page = int(gid), int(page)
    if page < 0 or not await db.one("SELECT 1 FROM groups WHERE id=?", (gid,)):
        return await c.answer("Sahifa yoki guruh topilmadi.", show_alert=True)
    await c.message.edit_reply_markup(reply_markup=await picker_kb(gid, page))
    await c.answer()


@r.callback_query(F.data.startswith("va:t:"), IsAdmin())
async def vid_toggle(c: CallbackQuery):
    _, _, gid, vid = c.data.split(":")
    gid, vid = int(gid), int(vid)
    if not await db.one("SELECT 1 FROM groups WHERE id=?", (gid,)) or not await db.one(
        "SELECT 1 FROM videos WHERE id=?", (vid,)
    ):
        return await c.answer("Guruh yoki video topilmadi.", show_alert=True)
    if await db.one("SELECT 1 FROM group_videos WHERE group_id=? AND video_id=?", (gid, vid)):
        await db.run("DELETE FROM group_videos WHERE group_id=? AND video_id=?", (gid, vid))
    else:
        await db.run(
            "INSERT INTO group_videos(group_id,video_id,position) VALUES(?,?,"
            "(SELECT COALESCE(MAX(position),0)+1 FROM group_videos WHERE group_id=?))",
            (gid, vid, gid),
        )
    current = 0
    for row in c.message.reply_markup.inline_keyboard:
        for button in row:
            if button.callback_data and button.callback_data.startswith("va:p:"):
                destination = int(button.callback_data.split(":")[3])
                current = destination + 1 if button.text.startswith("⬅️") else destination - 1
    await c.message.edit_reply_markup(reply_markup=await picker_kb(gid, current))
    await c.answer()


@r.callback_query(F.data.startswith("va:all:"), IsAdmin())
async def vid_all(c: CallbackQuery):
    parts = c.data.split(":")
    gid = int(parts[2])
    page = int(parts[3]) if len(parts) > 3 else 0
    if not await db.one("SELECT 1 FROM groups WHERE id=?", (gid,)):
        return await c.answer("Guruh topilmadi.", show_alert=True)
    for v in await db.fetch("SELECT id FROM videos ORDER BY id"):
        await db.run(
            "INSERT OR IGNORE INTO group_videos(group_id,video_id,position) VALUES(?,?,"
            "(SELECT COALESCE(MAX(position),0)+1 FROM group_videos WHERE group_id=?))",
            (gid, v["id"], gid),
        )
    await c.message.edit_reply_markup(reply_markup=await picker_kb(gid, page))
    await c.answer("Hammasi biriktirildi")


@r.callback_query(F.data == "va:done", IsAdmin())
async def vid_done(c: CallbackQuery):
    await c.message.edit_text("✅ Videolar biriktirildi.")
    await c.answer()


# ================================================================== ADMIN: TEST QO'SHISH
@r.message(F.text == A_TADD, IsAdmin())
async def test_add(m: Message, state: FSMContext):
    await state.clear()
    await state.set_state(TestAdd.title)
    await m.answer("Test nomini yozing (masalan: «Kvadrat tenglamalar»):")


@r.message(TestAdd.title, F.text)
async def test_add_title(m: Message, state: FSMContext):
    title = m.text.strip()
    if not title:
        return await m.answer("Test nomi bo'sh bo'lmasligi kerak.")
    await state.update_data(title=title)
    await state.set_state(TestAdd.content)
    await m.answer("Test materialini yuboring (PDF yoki rasm). Izoh ixtiyoriy:")


@r.message(TestAdd.content, IsAdmin())
async def test_add_content(m: Message, state: FSMContext):
    if m.photo:
        content = {"ctype": "photo", "file_id": m.photo[-1].file_id, "text": m.caption or ""}
    elif m.document:
        filename = (m.document.file_name or "").lower()
        if m.document.mime_type != "application/pdf" and not filename.endswith(".pdf"):
            return await m.answer("Hujjat PDF bo'lishi kerak. Rasm yuborishingiz ham mumkin.")
        content = {"ctype": "document", "file_id": m.document.file_id, "text": m.caption or ""}
    else:
        return await m.answer("Test uchun PDF fayl yoki rasm yuboring.")
    await state.update_data(**content)
    await state.set_state(TestAdd.code)
    await m.answer("Test uchun 4 xonali kod belgilang (masalan, 5239):")


@r.message(TestAdd.code, F.text, IsAdmin())
async def test_add_code(m: Message, state: FSMContext):
    code = m.text.strip()
    if not re.fullmatch(r"\d{4}", code):
        return await m.answer("Kod aynan 4 ta raqam bo'lishi kerak.")
    if await db.one("SELECT 1 FROM tests WHERE access_code=?", (code,)):
        return await m.answer("Bu kod boshqa testga berilgan. Boshqa 4 xonali kod kiriting.")
    await state.update_data(access_code=code)
    await state.set_state(TestAdd.confirm_code)
    await m.answer(
        f"Test kodi <code>{code}</code>. Shu kodni tasdiqlaysizmi?",
        reply_markup=ikb([
            [("✅ Tasdiqlash", "tac:confirm")],
            [("✏️ Boshqa kod kiritish", "tac:change")],
        ]),
    )


@r.callback_query(F.data == "tac:confirm", IsAdmin(), TestAdd.confirm_code)
async def test_add_code_confirm(c: CallbackQuery, state: FSMContext):
    await state.set_state(TestAdd.mode)
    await c.message.answer(
        "Test turini tanlang:",
        reply_markup=ikb([[("✍️ Qo'lda tekshirish", "tam:manual")], [("🤖 Avtomatik test", "tam:auto")]]),
    )
    await c.answer()


@r.callback_query(F.data == "tac:change", IsAdmin(), TestAdd.confirm_code)
async def test_add_code_change(c: CallbackQuery, state: FSMContext):
    await state.set_state(TestAdd.code)
    await c.message.answer("Yangi 4 xonali kodni kiriting:")
    await c.answer()


@r.callback_query(F.data.startswith("tam:"), IsAdmin(), TestAdd.mode)
async def test_add_mode(c: CallbackQuery, state: FSMContext):
    mode = c.data.split(":")[1]
    data = await state.get_data()
    if mode == "manual":
        await db.run(
            "INSERT INTO tests(title,ctype,file_id,text,mode,access_code) VALUES(?,?,?,?,?,?)",
            (data["title"], data["ctype"], data["file_id"], data["text"], "manual", data["access_code"]),
        )
        await state.clear()
        await c.message.answer(f"✅ Qo'lda tekshiriladigan test qo'shildi. Kod: <code>{data['access_code']}</code>")
    elif mode == "auto":
        await state.set_state(TestAdd.question_count)
        await c.message.answer("Avtomatik testdagi savollar sonini kiriting:")
    else:
        return await c.answer("Test turi noto'g'ri.", show_alert=True)
    await c.answer()


@r.message(TestAdd.question_count, F.text, IsAdmin())
async def test_add_question_count(m: Message, state: FSMContext):
    if not m.text.strip().isdigit() or int(m.text.strip()) < 1:
        return await m.answer("Savollar sonini 1 yoki undan katta butun son bilan kiriting.")
    total = int(m.text.strip())
    await state.update_data(question_count=total, questions=[])
    await state.set_state(TestAdd.auto_answer)
    await m.answer(
        "1-savolning to'g'ri javobini tanlang:",
        reply_markup=test_answer_keyboard(),
    )


def test_answer_keyboard(selected=None, prefix="tqa"):
    letters = [
        (f"{'✅ ' if selected == letter else ''}{letter}", f"{prefix}:{letter}")
        for letter in "ABCD"
    ]
    return ikb([letters, [("✅ Javobni tasdiqlash", f"{prefix}:confirm")]])


@r.callback_query(F.data.startswith("tqa:"), IsAdmin(), TestAdd.auto_answer)
async def test_add_auto_answer(c: CallbackQuery, state: FSMContext):
    choice = c.data.split(":")[1]
    if choice in {"A", "B", "C", "D"}:
        await state.update_data(pending_answer=choice)
        await c.message.edit_reply_markup(reply_markup=test_answer_keyboard(choice))
        return await c.answer(f"{choice} tanlandi")
    if choice != "confirm":
        return await c.answer("Javob tanlovi noto'g'ri.", show_alert=True)
    d = await state.get_data()
    correct = d.get("pending_answer")
    if not correct:
        return await c.answer("Avval to'g'ri javobni tanlang.", show_alert=True)
    questions = d.get("questions", [])
    questions.append({"correct": correct})
    await state.update_data(questions=questions, pending_answer=None)
    if len(questions) < d["question_count"]:
        await c.message.answer(
            f"✅ {len(questions)}-savol saqlandi.\n"
            f"{len(questions) + 1}-savolning to'g'ri javobini tanlang:",
            reply_markup=test_answer_keyboard(),
        )
    else:
        await c.message.answer(
            f"✅ {len(questions)} ta javob kaliti tayyor. Testni saqlang:",
            reply_markup=ikb([[("✅ Testni saqlash", "aq:finish")]]),
        )
    await c.answer()


@r.message(TestAdd.auto_answer)
async def test_add_auto_answer_invalid(m: Message):
    await m.answer("Javobni A, B, C yoki D tugmasi orqali tanlab, tasdiqlang.")


@r.callback_query(F.data == "aq:finish", IsAdmin(), TestAdd.auto_answer)
async def test_add_auto_finish(c: CallbackQuery, state: FSMContext):
    d = await state.get_data()
    questions = d.get("questions", [])
    if len(questions) != d.get("question_count"):
        return await c.answer(
            f"{d.get('question_count')} ta savoldan {len(questions)} tasi kiritilgan.",
            show_alert=True,
        )
    tid = await db.run(
        "INSERT INTO tests(title,ctype,file_id,text,mode,access_code) VALUES(?,?,?,?,?,?)",
        (d["title"], d["ctype"], d["file_id"], d["text"], "auto", d["access_code"]),
    )
    for position, question in enumerate(questions, 1):
        await db.run(
            "INSERT INTO test_questions(test_id,position,correct) VALUES(?,?,?)",
            (tid, position, question["correct"]),
        )
    await state.clear()
    await c.message.answer(
        f"✅ Avtomatik test qo'shildi: {esc(d['title'])} ({len(questions)} ta savol)\n"
        f"Test kodi: <code>{d['access_code']}</code>"
    )
    await c.answer()


@r.message(F.text == A_TDEL, IsAdmin())
async def test_delete_menu(m: Message):
    rows = await db.fetch("SELECT id,title,mode FROM tests ORDER BY id")
    if not rows:
        return await m.answer("O'chiriladigan test yo'q.")
    await m.answer(
        "O'chirmoqchi bo'lgan testni tanlang:",
        reply_markup=ikb([
            [(f"{'🤖' if row['mode'] == 'auto' else '✍️'} {row['title']}"[:60], f"tdel:ask:{row['id']}")]
            for row in rows
        ]),
    )


@r.callback_query(F.data.startswith("tdel:ask:"), IsAdmin())
async def test_delete_confirm(c: CallbackQuery):
    tid = int(c.data.split(":")[2])
    test = await db.one("SELECT title FROM tests WHERE id=?", (tid,))
    if not test:
        return await c.answer("Test topilmadi.", show_alert=True)
    await c.message.answer(
        f"«{esc(test['title'])}» testini va uning natijalarini butunlay o'chirasizmi?",
        reply_markup=ikb([[("🗑 Ha, o'chirish", f"tdel:yes:{tid}")], [("Bekor qilish", "tdel:no")]]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("tdel:yes:"), IsAdmin())
async def test_delete(c: CallbackQuery):
    tid = int(c.data.split(":")[2])
    test = await db.one("SELECT id FROM tests WHERE id=?", (tid,))
    if not test:
        return await c.answer("Test topilmadi.", show_alert=True)
    await db.run(
        "DELETE FROM sub_files WHERE sub_id IN (SELECT id FROM submissions WHERE test_id=?)", (tid,)
    )
    await db.run("DELETE FROM submissions WHERE test_id=?", (tid,))
    await db.run("DELETE FROM attempts WHERE test_id=?", (tid,))
    await db.run("DELETE FROM test_questions WHERE test_id=?", (tid,))
    await db.run("DELETE FROM tests WHERE id=?", (tid,))
    await c.message.edit_text("✅ Test va unga tegishli natijalar o'chirildi.")
    await c.answer()


@r.callback_query(F.data == "tdel:no", IsAdmin())
async def test_delete_cancel(c: CallbackQuery):
    await c.message.edit_text("O'chirish bekor qilindi.")
    await c.answer()


@r.message(F.text == A_TEDIT, IsAdmin())
async def test_edit_menu(m: Message, state: FSMContext):
    await state.clear()
    tests = await db.fetch("SELECT id, title, mode, access_code FROM tests ORDER BY id DESC")
    if not tests:
        return await m.answer("Tahrirlanadigan test yo'q.")
    await m.answer(
        "Tahrirlamoqchi bo'lgan testni tanlang:",
        reply_markup=ikb([
            [(
                f"{'🤖' if test['mode'] == 'auto' else '✍️'} {test['title']} "
                f"({test['access_code'] or 'eski'})"[:60],
                f"tedit:pick:{test['id']}",
            )]
            for test in tests
        ]),
    )


@r.callback_query(F.data.startswith("tedit:pick:"), IsAdmin())
async def test_edit_pick(c: CallbackQuery):
    test_id = int(c.data.split(":")[2])
    test = await db.one("SELECT id, title, mode FROM tests WHERE id=?", (test_id,))
    if not test:
        return await c.answer("Test topilmadi.", show_alert=True)
    content_controls = (
        [
            [("📄 Test materialini almashtirish", f"tedit:material:{test_id}")],
            [("✅ Javob kalitlarini almashtirish", f"tedit:content:{test_id}")],
        ]
        if test["mode"] == "auto"
        else [[("📄 Materialni almashtirish", f"tedit:content:{test_id}")]]
    )
    await c.message.answer(
        f"«{esc(test['title'])}» testini qanday tahrirlaymiz?",
        reply_markup=ikb([
            [("✏️ Nomini o'zgartirish", f"tedit:title:{test_id}")],
            [("🔑 Kirish kodini o'zgartirish", f"tedit:code:{test_id}")],
        ] + content_controls),
    )
    await c.answer()


@r.callback_query(F.data.startswith("tedit:title:"), IsAdmin())
async def test_edit_title_start(c: CallbackQuery, state: FSMContext):
    test_id = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM tests WHERE id=?", (test_id,)):
        return await c.answer("Test topilmadi.", show_alert=True)
    await state.set_state(TestEdit.title)
    await state.update_data(test_id=test_id)
    await c.message.answer("Testning yangi nomini yozing:")
    await c.answer()


@r.message(TestEdit.title, F.text, IsAdmin())
async def test_edit_title(m: Message, state: FSMContext):
    title = m.text.strip()
    if not title:
        return await m.answer("Test nomi bo'sh bo'lmasligi kerak.")
    data = await state.get_data()
    await db.run("UPDATE tests SET title=? WHERE id=?", (title, data["test_id"]))
    await state.clear()
    await m.answer("✅ Test nomi o'zgartirildi.")


@r.callback_query(F.data.startswith("tedit:code:"), IsAdmin())
async def test_edit_code_start(c: CallbackQuery, state: FSMContext):
    test_id = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM tests WHERE id=?", (test_id,)):
        return await c.answer("Test topilmadi.", show_alert=True)
    await state.set_state(TestEdit.code)
    await state.update_data(test_id=test_id)
    await c.message.answer("Test uchun yangi 4 xonali kodni kiriting:")
    await c.answer()


@r.message(TestEdit.code, F.text, IsAdmin())
async def test_edit_code(m: Message, state: FSMContext):
    code = m.text.strip()
    if not re.fullmatch(r"\d{4}", code):
        return await m.answer("Kod aynan 4 ta raqam bo'lishi kerak.")
    data = await state.get_data()
    if await db.one(
        "SELECT 1 FROM tests WHERE access_code=? AND id!=?",
        (code, data["test_id"]),
    ):
        return await m.answer("Bu kod boshqa testga berilgan. Boshqa kod kiriting.")
    await db.run("UPDATE tests SET access_code=? WHERE id=?", (code, data["test_id"]))
    await state.clear()
    await m.answer(f"✅ Testning yangi kodi: <code>{code}</code>")


@r.callback_query(F.data.startswith("tedit:content:"), IsAdmin())
async def test_edit_content_start(c: CallbackQuery, state: FSMContext):
    test_id = int(c.data.split(":")[2])
    test = await db.one("SELECT mode FROM tests WHERE id=?", (test_id,))
    if not test:
        return await c.answer("Test topilmadi.", show_alert=True)
    await state.update_data(test_id=test_id)
    if test["mode"] == "auto":
        await state.update_data(questions=[])
        await state.set_state(TestEdit.auto_count)
        await c.message.answer("Yangi javob kalitida nechta savol bo'lishini kiriting:")
    else:
        await state.set_state(TestEdit.manual_content)
        await c.message.answer("Yangi test materialini yuboring: matn, rasm yoki fayl (PDF).")
    await c.answer()


@r.callback_query(F.data.startswith("tedit:material:"), IsAdmin())
async def test_edit_material_start(c: CallbackQuery, state: FSMContext):
    test_id = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM tests WHERE id=?", (test_id,)):
        return await c.answer("Test topilmadi.", show_alert=True)
    await state.set_state(TestEdit.manual_content)
    await state.update_data(test_id=test_id)
    await c.message.answer("Yangi test materialini yuboring (PDF yoki rasm).")
    await c.answer()


@r.message(TestEdit.auto_count, F.text, IsAdmin())
async def test_edit_auto_count(m: Message, state: FSMContext):
    if not m.text.strip().isdigit() or int(m.text.strip()) < 1:
        return await m.answer("Savollar sonini 1 yoki undan katta butun son bilan kiriting.")
    total = int(m.text.strip())
    await state.update_data(question_count=total, questions=[])
    await state.set_state(TestEdit.auto_answer)
    await m.answer(
        "1-savolning to'g'ri javobini tanlang:",
        reply_markup=test_answer_keyboard(prefix="tqe"),
    )


@r.message(TestEdit.manual_content, IsAdmin())
async def test_edit_manual_content(m: Message, state: FSMContext):
    if m.photo:
        ctype, file_id, text = "photo", m.photo[-1].file_id, m.caption
    elif m.document:
        ctype, file_id, text = "document", m.document.file_id, m.caption
    elif m.text:
        ctype, file_id, text = "text", None, m.text
    else:
        return await m.answer("Matn, rasm yoki fayl yuboring.")
    data = await state.get_data()
    await db.run(
        "UPDATE tests SET ctype=?, file_id=?, text=? WHERE id=?",
        (ctype, file_id, text, data["test_id"]),
    )
    await state.clear()
    await m.answer("✅ Test materiali almashtirildi.")


@r.callback_query(F.data.startswith("tqe:"), IsAdmin(), TestEdit.auto_answer)
async def test_edit_auto_answer(c: CallbackQuery, state: FSMContext):
    choice = c.data.split(":")[1]
    if choice in {"A", "B", "C", "D"}:
        await state.update_data(pending_answer=choice)
        await c.message.edit_reply_markup(
            reply_markup=test_answer_keyboard(choice, prefix="tqe")
        )
        return await c.answer(f"{choice} tanlandi")
    if choice != "confirm":
        return await c.answer("Javob tanlovi noto'g'ri.", show_alert=True)
    data = await state.get_data()
    correct = data.get("pending_answer")
    if not correct:
        return await c.answer("Avval to'g'ri javobni tanlang.", show_alert=True)
    questions = data.get("questions", [])
    questions.append({"correct": correct})
    await state.update_data(questions=questions, pending_answer=None)
    if len(questions) < data["question_count"]:
        await c.message.answer(
            f"✅ {len(questions)}-savol saqlandi.\n"
            f"{len(questions) + 1}-savolning to'g'ri javobini tanlang:",
            reply_markup=test_answer_keyboard(prefix="tqe"),
        )
    else:
        await c.message.answer(
            f"✅ {len(questions)} ta javob kaliti tayyor.",
            reply_markup=ikb([[("✅ O'zgarishlarni saqlash", "tedit:q:finish")]]),
        )
    await c.answer()


@r.message(TestEdit.auto_answer, IsAdmin())
async def test_edit_auto_answer_invalid(m: Message):
    await m.answer("Javobni A, B, C yoki D tugmasi orqali tanlab, tasdiqlang.")


@r.callback_query(F.data == "tedit:q:finish", IsAdmin(), TestEdit.auto_answer)
async def test_edit_auto_finish(c: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    questions = data.get("questions", [])
    if len(questions) != data.get("question_count"):
        return await c.answer(
            f"{data.get('question_count')} ta javobdan {len(questions)} tasi tayyor.",
            show_alert=True,
        )
    test_id = data["test_id"]
    if not await db.one("SELECT 1 FROM tests WHERE id=?", (test_id,)):
        await state.clear()
        return await c.answer("Test topilmadi.", show_alert=True)
    await db.run("DELETE FROM test_questions WHERE test_id=?", (test_id,))
    for position, question in enumerate(questions, 1):
        await db.run(
            "INSERT INTO test_questions(test_id,position,correct) VALUES(?,?,?)",
            (test_id, position, question["correct"]),
        )
    await state.clear()
    await c.message.answer(f"✅ Test savollari almashtirildi ({len(questions)} ta).")
    await c.answer()


# ================================================================== ADMIN: KITOB QO'SHISH
@r.message(F.text == A_BADD, IsAdmin())
async def book_add(m: Message, state: FSMContext):
    await state.clear()
    await m.answer(
        "Kitob qaysi fandan?",
        reply_markup=ikb([[("📐 Algebra", "bka:Algebra"), ("📏 Geometriya", "bka:Geometriya")], [("📚 Boshqa", "bka:Boshqa")]]),
    )


@r.callback_query(F.data.startswith("bka:"), IsAdmin())
async def book_add_subject(c: CallbackQuery, state: FSMContext):
    await state.set_state(BookAdd.file)
    await state.update_data(subject=c.data.split(":")[1])
    await c.message.answer("Kitob faylini (PDF) yuboring. Caption yozsangiz, o'sha nom bo'ladi.")
    await c.answer()


@r.message(BookAdd.file, F.document)
async def book_add_file(m: Message, state: FSMContext):
    d = await state.get_data()
    title = (m.caption or m.document.file_name or "Kitob").strip()
    await db.run("INSERT INTO books(subject,title,file_id) VALUES(?,?,?)", (d["subject"], title, m.document.file_id))
    await state.clear()
    await m.answer(f"✅ Kitob qo'shildi: {esc(title)}")


@r.message(F.text == A_BEDIT, IsAdmin())
async def book_edit_menu(m: Message, state: FSMContext):
    await state.clear()
    books = await db.fetch("SELECT id, subject, title FROM books ORDER BY subject, title")
    if not books:
        return await m.answer("Tahrirlanadigan kitob yo'q.")
    await m.answer(
        "Tahrirlamoqchi bo'lgan kitobni tanlang:",
        reply_markup=ikb([
            [(f"{book['subject']} — {book['title']}"[:60], f"bkedit:pick:{book['id']}")]
            for book in books
        ]),
    )


@r.callback_query(F.data.startswith("bkedit:pick:"), IsAdmin())
async def book_edit_pick(c: CallbackQuery):
    book_id = int(c.data.split(":")[2])
    book = await db.one("SELECT title FROM books WHERE id=?", (book_id,))
    if not book:
        return await c.answer("Kitob topilmadi.", show_alert=True)
    await c.message.answer(
        f"«{esc(book['title'])}» kitobini qanday tahrirlaymiz?",
        reply_markup=ikb([
            [("✏️ Fan va nom", f"bkedit:details:{book_id}")],
            [("📄 Faylni almashtirish", f"bkedit:file:{book_id}")],
        ]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("bkedit:details:"), IsAdmin())
async def book_edit_details_start(c: CallbackQuery, state: FSMContext):
    book_id = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM books WHERE id=?", (book_id,)):
        return await c.answer("Kitob topilmadi.", show_alert=True)
    await state.set_state(BookEdit.details)
    await state.update_data(book_id=book_id)
    await c.message.answer("Yangi fan va nomni shu ko'rinishda yozing:\nFan | Kitob nomi")
    await c.answer()


@r.message(BookEdit.details, F.text, IsAdmin())
async def book_edit_details(m: Message, state: FSMContext):
    if "|" not in m.text:
        return await m.answer("Fan va nom orasiga | belgisi qo'ying. Masalan: Algebra | Tenglamalar")
    subject, title = (part.strip() for part in m.text.split("|", 1))
    if not subject or not title:
        return await m.answer("Fan ham, kitob nomi ham bo'sh bo'lmasligi kerak.")
    subject = next((known for known in SUBJECTS if known.casefold() == subject.casefold()), None)
    if not subject:
        return await m.answer("Fan Algebra, Geometriya yoki Boshqa bo'lishi kerak.")
    data = await state.get_data()
    await db.run("UPDATE books SET subject=?, title=? WHERE id=?", (subject, title, data["book_id"]))
    await state.clear()
    await m.answer("✅ Kitob ma'lumoti o'zgartirildi.")


@r.callback_query(F.data.startswith("bkedit:file:"), IsAdmin())
async def book_edit_file_start(c: CallbackQuery, state: FSMContext):
    book_id = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM books WHERE id=?", (book_id,)):
        return await c.answer("Kitob topilmadi.", show_alert=True)
    await state.set_state(BookEdit.file)
    await state.update_data(book_id=book_id)
    await c.message.answer("Yangi PDF faylini yuboring:")
    await c.answer()


@r.message(BookEdit.file, F.document, IsAdmin())
async def book_edit_file(m: Message, state: FSMContext):
    data = await state.get_data()
    await db.run("UPDATE books SET file_id=? WHERE id=?", (m.document.file_id, data["book_id"]))
    await state.clear()
    await m.answer("✅ Kitob fayli almashtirildi.")


@r.message(F.text == A_BDEL, IsAdmin())
async def book_delete_menu(m: Message, state: FSMContext):
    await state.clear()
    rows = await db.fetch("SELECT id, subject, title FROM books ORDER BY subject, title")
    if not rows:
        return await m.answer("O'chiriladigan kitob yo'q.")
    await m.answer(
        "O'chirmoqchi bo'lgan kitobni tanlang:",
        reply_markup=ikb([
            [(f"{row['subject']} — {row['title']}"[:60], f"bkdel:ask:{row['id']}")]
            for row in rows
        ]),
    )


@r.callback_query(F.data.startswith("bkdel:ask:"), IsAdmin())
async def book_delete_confirm(c: CallbackQuery):
    book_id = int(c.data.split(":")[2])
    book = await db.one("SELECT title FROM books WHERE id=?", (book_id,))
    if not book:
        return await c.answer("Kitob topilmadi.", show_alert=True)
    await c.message.answer(
        f"«{esc(book['title'])}» kitobini o'chirasizmi?",
        reply_markup=ikb([
            [("🗑 Ha, o'chirish", f"bkdel:yes:{book_id}")],
            [("Bekor qilish", "bkdel:no")],
        ]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("bkdel:yes:"), IsAdmin())
async def book_delete(c: CallbackQuery):
    book_id = int(c.data.split(":")[2])
    if not await db.one("SELECT 1 FROM books WHERE id=?", (book_id,)):
        return await c.answer("Kitob topilmadi.", show_alert=True)
    await db.run("DELETE FROM books WHERE id=?", (book_id,))
    await c.message.edit_text("✅ Kitob o'chirildi.")
    await c.answer()


@r.callback_query(F.data == "bkdel:no", IsAdmin())
async def book_delete_cancel(c: CallbackQuery):
    await c.message.edit_text("O'chirish bekor qilindi.")
    await c.answer()


# ================================================================== ADMIN: TEST NATIJALARI
@r.message(F.text == A_RES, IsAdmin())
async def results(m: Message, state: FSMContext):
    await state.clear()
    rows = await db.fetch("SELECT id,name FROM groups ORDER BY id")
    if not rows:
        return await m.answer("Guruhlar yo'q.")
    await m.answer("Guruhni tanlang:", reply_markup=ikb([[(x["name"], f"tr:g:{x['id']}")] for x in rows]))


@r.callback_query(F.data.startswith("tr:g:"), IsAdmin())
async def results_group(c: CallbackQuery):
    gid = int(c.data.split(":")[2])
    tests = await db.fetch("SELECT id,title FROM tests ORDER BY id")
    if not tests:
        return await c.answer("Testlar yo'q.", show_alert=True)
    await c.message.answer("Testni tanlang:", reply_markup=ikb([[(t["title"][:50], f"tr:t:{gid}:{t['id']}")] for t in tests]))
    await c.answer()


@r.callback_query(F.data.startswith("tr:t:"), IsAdmin())
async def results_test(c: CallbackQuery):
    _, _, gid, tid = c.data.split(":")
    gid, tid = int(gid), int(tid)
    t = await db.one("SELECT title,mode FROM tests WHERE id=?", (tid,))
    if not t:
        return await c.answer("Test topilmadi.", show_alert=True)
    g = await db.one("SELECT name FROM groups WHERE id=?", (gid,))
    studs = await db.fetch("SELECT id, full_name FROM users WHERE group_id=? ORDER BY full_name", (gid,))
    icons = {"approved": "✅", "pending": "⏳", "rejected": "❌"}
    lines, cnt = [], {"approved": 0, "pending": 0, "rejected": 0, None: 0}
    for s in studs:
        if t["mode"] == "auto":
            attempt = await db.one(
                "SELECT score,total,finished_at FROM attempts WHERE test_id=? AND user_id=? "
                "ORDER BY id DESC LIMIT 1",
                (tid, s["id"]),
            )
            st = "approved" if attempt and attempt["finished_at"] else None
            score_text = (
                f" — {attempt['score']}/{attempt['total']}"
                if attempt and attempt["finished_at"] else ""
            )
        else:
            sub = await db.one(
                "SELECT status FROM submissions WHERE test_id=? AND user_id=? ORDER BY id DESC LIMIT 1",
                (tid, s["id"]),
            )
            st = sub["status"] if sub else None
            score_text = ""
        cnt[st] += 1
        lines.append(f"{icons.get(st, '➖')} {esc(s['full_name'])}{score_text}")
    head = (
        f"🧾 <b>{esc(t['title'])}</b> — {esc(g['name'])}\n"
        f"✅ {cnt['approved']}  ⏳ {cnt['pending']}  ❌ {cnt['rejected']}  ➖ {cnt[None]}\n"
        "(➖ topshirmagan)\n\n"
    )
    await c.message.answer(head + ("\n".join(lines) or "O'quvchilar yo'q."))
    await c.answer()


# ================================================================== ADMIN: XABAR YUBORISH
@r.message(F.text == A_BC, IsAdmin())
async def bc_start(m: Message, state: FSMContext):
    await state.clear()
    await m.answer(
        "Kimga yuboramiz?",
        reply_markup=ikb([[("👥 Hammaga", "bc:all")], [("🏷 Guruhga", "bc:grp")], [("👤 Bitta o'quvchiga", "bc:one")]]),
    )


async def bc_ask(c: CallbackQuery, state: FSMContext, target):
    await state.set_state(Bc.msg)
    await state.update_data(target=target)
    await c.message.answer("Xabarni yuboring (matn, rasm, video, ovozli xabar — hammasi mumkin):")
    await c.answer()


@r.callback_query(F.data == "bc:all", IsAdmin())
async def bc_all(c: CallbackQuery, state: FSMContext):
    await bc_ask(c, state, "all")


@r.callback_query(F.data == "bc:grp", IsAdmin())
async def bc_grp(c: CallbackQuery):
    rows = await db.fetch("SELECT id,name FROM groups ORDER BY id")
    if not rows:
        return await c.answer("Guruhlar yo'q.", show_alert=True)
    await c.message.answer("Guruhni tanlang:", reply_markup=ikb([[(x["name"], f"bc:g:{x['id']}")] for x in rows]))
    await c.answer()


@r.callback_query(F.data.startswith("bc:g:"), IsAdmin())
async def bc_grp_pick(c: CallbackQuery, state: FSMContext):
    await bc_ask(c, state, "g:" + c.data.split(":")[2])


@r.callback_query(F.data == "bc:one", IsAdmin())
async def bc_one(c: CallbackQuery):
    studs = await db.fetch("SELECT id, full_name FROM users ORDER BY full_name LIMIT 80")
    if not studs:
        return await c.answer("O'quvchilar yo'q.", show_alert=True)
    await c.message.answer(
        "O'quvchini tanlang:", reply_markup=ikb([[(s["full_name"][:50], f"bc:u:{s['id']}")] for s in studs])
    )
    await c.answer()


@r.callback_query(F.data.startswith("bc:u:"), IsAdmin())
async def bc_one_pick(c: CallbackQuery, state: FSMContext):
    await bc_ask(c, state, "u:" + c.data.split(":")[2])


@r.message(Bc.msg)
async def bc_send(m: Message, state: FSMContext):
    d = await state.get_data()
    t = d["target"]
    await state.clear()
    if t == "all":
        rows = await db.fetch("SELECT id FROM users WHERE blocked=0")
    elif t.startswith("g:"):
        rows = await db.fetch("SELECT id FROM users WHERE group_id=? AND blocked=0", (int(t[2:]),))
    else:
        rows = await db.fetch("SELECT id FROM users WHERE id=?", (int(t[2:]),))
    ok = fail = 0
    for row in rows:
        try:
            await bot.copy_message(row["id"], m.chat.id, m.message_id)
            ok += 1
        except TelegramForbiddenError:
            await db.run("UPDATE users SET blocked=1 WHERE id=?", (row["id"],))
            fail += 1
        except Exception:
            fail += 1
        await asyncio.sleep(0.05)
    await m.answer(f"📢 Yuborildi: {ok} ta\n⚠️ Yetmadi: {fail} ta")


# ================================================================== ADMIN: STATISTIKA
@r.message(F.text == A_STAT, IsAdmin())
async def stats(m: Message, state: FSMContext):
    await state.clear()
    now = int(time.time())
    total = await db.count("SELECT COUNT(*) FROM users")
    in_group = await db.count("SELECT COUNT(*) FROM users WHERE group_id IS NOT NULL AND blocked=0")
    left = await db.count("SELECT COUNT(*) FROM users WHERE blocked=1")
    added24 = await db.count("SELECT COUNT(*) FROM users WHERE created_at>=?", (now - 86400,))
    added7 = await db.count("SELECT COUNT(*) FROM users WHERE created_at>=?", (now - 7 * 86400,))
    added30 = await db.count("SELECT COUNT(*) FROM users WHERE created_at>=?", (now - 30 * 86400,))
    a24 = await db.count("SELECT COUNT(*) FROM users WHERE last_active>=?", (now - 86400,))
    a7 = await db.count("SELECT COUNT(*) FROM users WHERE last_active>=?", (now - 7 * 86400,))
    a30 = await db.count("SELECT COUNT(*) FROM users WHERE last_active>=?", (now - 30 * 86400,))
    buttons = await db.count(
        "SELECT COUNT(DISTINCT CASE WHEN instr(action,':')>0 "
        "THEN substr(action,1,instr(action,':')-1) ELSE action END) "
        "FROM user_activity WHERE event_type='callback'"
    )
    clicks = await db.count("SELECT COUNT(*) FROM user_activity WHERE event_type='callback'")
    content = sum([await db.count(f"SELECT COUNT(*) FROM {table}") for table in ("videos", "books", "tests")])
    me = await bot.get_me()
    checked_at = time.strftime("%Y-%m-%d %H:%M", time.localtime(now))
    await m.answer(
        "📊 <b>Statistika</b>\n\n"
        f"• Obunachilar soni: {total} ta\n"
        f"• Faol obunachilar: {in_group} ta\n"
        f"• Tark etganlar: {left} ta\n\n"
        "📈 <b>Obunachilar qo'shilishi</b>\n"
        f"• Oxirgi 24 soat: +{added24} obunachi\n"
        f"• Oxirgi 7 kun: +{added7} obunachi\n"
        f"• Oxirgi 30 kun: +{added30} obunachi\n\n"
        "📊 <b>Faollik</b>\n"
        f"• Oxirgi 24 soatda faol: {a24} ta\n"
        f"• Oxirgi 7 kunda faol: {a7} ta\n"
        f"• Oxirgi 30 kunda faol: {a30} ta\n\n"
        "🧩 <b>Menyu</b>\n"
        f"• Tugmalar: {buttons} ta\n"
        f"• Kontentlar: {content} ta\n"
        f"• Jami bosilgan: {clicks} marta\n\n"
        f"🕒 {checked_at} holatiga\n"
        f"🤖 @{esc(me.username or 'bot')}"
    )


@r.message(F.text == A_REPORT, IsAdmin())
async def report_menu(m: Message, state: FSMContext):
    await state.clear()
    groups = await db.fetch("SELECT id,name FROM groups ORDER BY name")
    if not groups:
        return await m.answer("Hisobot uchun guruhlar mavjud emas.")
    await m.answer(
        "Excel hisobotini olish uchun guruhni tanlang:",
        reply_markup=ikb([[(group["name"], f"report:g:{group['id']}")] for group in groups]),
    )


def excel_text(value):
    text = "" if value is None else str(value)
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


@r.callback_query(F.data.startswith("report:g:"), IsAdmin())
async def report_group(c: CallbackQuery):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    group_id = int(c.data.split(":")[2])
    group = await db.one("SELECT name FROM groups WHERE id=?", (group_id,))
    if not group:
        return await c.answer("Guruh topilmadi.", show_alert=True)
    students = await db.fetch(
        "SELECT u.id,u.full_name,u.phone,u.username,u.created_at,u.last_active,u.blocked "
        "FROM users u WHERE u.group_id=? ORDER BY u.full_name,u.id",
        (group_id,),
    )
    activity = await db.fetch(
        "SELECT a.user_id,u.full_name,u.phone,a.description,a.category,a.event_type,a.occurred_at "
        "FROM user_activity a JOIN users u ON u.id=a.user_id "
        "WHERE u.group_id=? ORDER BY a.occurred_at,a.id",
        (group_id,),
    )
    counts = {}
    summaries = {}
    for event in activity:
        uid = event["user_id"]
        counts[uid] = counts.get(uid, 0) + 1
        summaries.setdefault(uid, []).append(event["description"])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Oquvchilar"
    sheet.append([
        "Telegram ID", "Ism-familiya", "Telefon", "Username", "Guruh",
        "Ro'yxatdan o'tgan", "Oxirgi faollik", "Botdan foydalanish soni",
        "Bajargan amallari", "Holati",
    ])
    for column in range(1, 11):
        sheet.cell(row=1, column=column).font = Font(bold=True, color="FFFFFF")
        sheet.cell(row=1, column=column).fill = PatternFill("solid", fgColor="305496")
    for student in students:
        sheet.append([
            student["id"], excel_text(student["full_name"]), excel_text(student["phone"]),
            excel_text(f"@{student['username']}" if student["username"] else ""),
            excel_text(group["name"]),
            time.strftime("%Y-%m-%d %H:%M", time.localtime(student["created_at"]))
            if student["created_at"] else "",
            time.strftime("%Y-%m-%d %H:%M", time.localtime(student["last_active"]))
            if student["last_active"] else "",
            counts.get(student["id"], 0),
            excel_text("; ".join(dict.fromkeys(summaries.get(student["id"], [])))[:32000]),
            "Botni bloklagan" if student["blocked"] else "Faol",
        ])
    for column, width in enumerate((16, 26, 20, 22, 24, 22, 22, 22, 60, 20), 1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    log = workbook.create_sheet("Faollik tarixi")
    log.append(["Telegram ID", "Ism-familiya", "Telefon", "Amal", "Bo'lim", "Turi", "Sana-vaqt"])
    for column in range(1, 8):
        log.cell(row=1, column=column).font = Font(bold=True, color="FFFFFF")
        log.cell(row=1, column=column).fill = PatternFill("solid", fgColor="305496")
    for event in activity:
        log.append([
            event["user_id"], excel_text(event["full_name"]), excel_text(event["phone"]),
            excel_text(event["description"]), excel_text(event["category"]), excel_text(event["event_type"]),
            time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(event["occurred_at"])),
        ])
    for column, width in enumerate((16, 26, 20, 50, 16, 16, 24), 1):
        log.column_dimensions[chr(64 + column)].width = width
    log.freeze_panes = "A2"
    log.auto_filter.ref = log.dimensions

    output = BytesIO()
    workbook.save(output)
    filename = f"hisobot_{re.sub(r'[^A-Za-z0-9_-]+', '_', group['name'])}.xlsx"
    await bot.send_document(
        c.from_user.id,
        BufferedInputFile(output.getvalue(), filename=filename),
        caption=f"📑 «{esc(group['name'])}» guruhi bo'yicha Excel hisobot.",
    )
    await c.answer("Hisobot yuborildi.")


async def reminder_loop():
    while True:
        await asyncio.sleep(3600)
        now = int(time.time())
        users = await db.fetch(
            "SELECT u.id FROM users u WHERE u.group_id IS NOT NULL AND u.blocked=0 "
            "AND COALESCE(u.last_reminded_at,0)<=? "
            "AND NOT EXISTS (SELECT 1 FROM user_activity a WHERE a.user_id=u.id "
            "AND a.category IN ('video','test') AND a.occurred_at>=?)",
            (now - 3600, now - 86400),
        )
        for user in users:
            sent = await safe_send(
                user["id"],
                "📚 Darslarni davom ettirish vaqti! Video darslarni ko'ring yoki MOCK test ishlang.",
            )
            if sent:
                await db.run("UPDATE users SET last_reminded_at=? WHERE id=?", (now, user["id"]))
            await asyncio.sleep(0.05)


# ================================================================== ADMIN: ADMINLAR
@r.message(F.text == A_ADMINS, IsAdmin())
async def admins_list(m: Message, state: FSMContext):
    await state.clear()
    rows = await db.fetch("SELECT id, full_name FROM admins")
    lines = []
    controls = []
    for admin in rows:
        lines.append(f"• {esc(admin['full_name'])} — <code>{admin['id']}</code>")
        controls.append([
            (f"✏️ {str(admin['full_name'] or 'Admin')[:35]}", f"ad:edit:{admin['id']}"),
            ("🗑 O'chirish", f"ad:del:ask:{admin['id']}"),
        ])
    controls.append([("➕ Admin qo'shish", "ad:new")])
    text = "👮 <b>Adminlar</b>\n\n" + ("\n".join(lines) or "Adminlar yo'q.")
    await m.answer(text, reply_markup=ikb(controls))


@r.callback_query(F.data == "ad:new", IsAdmin())
async def admin_new(c: CallbackQuery, state: FSMContext):
    await state.set_state(AdmAdd.name)
    await c.message.answer("Yangi adminning ism-familiyasini yozing:")
    await c.answer()


@r.message(AdmAdd.name, F.text, IsAdmin())
async def admin_new_name(m: Message, state: FSMContext):
    name = m.text.strip()
    if not name:
        return await m.answer("Ism bo'sh bo'lmasligi kerak.")
    await state.update_data(name=name)
    await state.set_state(AdmAdd.uid)
    await m.answer("Uning Telegram ID raqamini yozing (raqam ko'rinishida):")


@r.message(AdmAdd.uid, F.text, IsAdmin())
async def admin_new_id(m: Message, state: FSMContext):
    if not m.text.strip().isdigit():
        return await m.answer("ID faqat raqamlardan iborat bo'lishi kerak.")
    d = await state.get_data()
    uid = int(m.text.strip())
    if await db.one("SELECT 1 FROM admins WHERE id=?", (uid,)):
        await state.clear()
        return await m.answer("Bu foydalanuvchi allaqachon admin.")
    await db.run("INSERT OR IGNORE INTO admins(id,full_name) VALUES(?,?)", (uid, d["name"]))
    if not await db.one("SELECT 1 FROM admins WHERE id=?", (uid,)):
        await state.clear()
        return await m.answer("Adminni saqlab bo'lmadi.")
    await state.clear()
    await m.answer(f"✅ {esc(d['name'])} admin qilindi.")
    await safe_send(uid, "🎉 Siz admin etib tayinlandingiz. /start bosing.")


@r.callback_query(F.data.startswith("ad:edit:"), IsAdmin())
async def admin_edit_start(c: CallbackQuery, state: FSMContext):
    uid = int(c.data.split(":")[2])
    admin = await db.one("SELECT full_name FROM admins WHERE id=?", (uid,))
    if not admin:
        return await c.answer("Admin topilmadi.", show_alert=True)
    await state.set_state(AdmEdit.name)
    await state.update_data(admin_id=uid)
    await c.message.answer(f"Adminning yangi ismini yozing (Telegram ID: {uid}):")
    await c.answer()


@r.message(AdmEdit.name, F.text, IsAdmin())
async def admin_edit_name(m: Message, state: FSMContext):
    name = m.text.strip()
    if not name:
        return await m.answer("Ism bo'sh bo'lmasligi kerak.")
    data = await state.get_data()
    await db.run("UPDATE admins SET full_name=? WHERE id=?", (name, data["admin_id"]))
    await state.clear()
    await m.answer("✅ Admin ma'lumoti o'zgartirildi.")


@r.callback_query(F.data.startswith("ad:del:ask:"), IsAdmin())
async def admin_delete_confirm(c: CallbackQuery):
    uid = int(c.data.split(":")[3])
    admin = await db.one("SELECT full_name FROM admins WHERE id=?", (uid,))
    if not admin:
        return await c.answer("Admin topilmadi.", show_alert=True)
    if uid == c.from_user.id:
        return await c.answer("O'zingizni adminlikdan o'chira olmaysiz.", show_alert=True)
    if uid in ADMIN_IDS:
        return await c.answer("Asosiy adminni o'chirib bo'lmaydi.", show_alert=True)
    await c.message.answer(
        f"«{esc(admin['full_name'])}» adminini o'chirasizmi?",
        reply_markup=ikb([
            [("🗑 Ha, o'chirish", f"ad:del:yes:{uid}")],
            [("Bekor qilish", "ad:del:no")],
        ]),
    )
    await c.answer()


@r.callback_query(F.data.startswith("ad:del:yes:"), IsAdmin())
async def admin_delete(c: CallbackQuery):
    uid = int(c.data.split(":")[3])
    if uid == c.from_user.id:
        return await c.answer("O'zingizni adminlikdan o'chira olmaysiz.", show_alert=True)
    if uid in ADMIN_IDS:
        return await c.answer("Asosiy adminni o'chirib bo'lmaydi.", show_alert=True)
    admin = await db.one("SELECT 1 FROM admins WHERE id=?", (uid,))
    if not admin:
        return await c.answer("Admin topilmadi.", show_alert=True)
    if await db.count("SELECT COUNT(*) FROM admins") <= 1:
        return await c.answer("Oxirgi adminni o'chirib bo'lmaydi.", show_alert=True)
    await db.run("DELETE FROM admins WHERE id=?", (uid,))
    await safe_send(uid, "Sizning adminlik huquqingiz olib tashlandi.")
    await c.message.edit_text("✅ Admin o'chirildi.")
    await c.answer()


@r.callback_query(F.data == "ad:del:no", IsAdmin())
async def admin_delete_cancel(c: CallbackQuery):
    await c.message.edit_text("Adminni o'chirish bekor qilindi.")
    await c.answer()


# ================================================================== XATOLIKLAR
@dp.error()
async def on_error(event: ErrorEvent):
    logging.exception("Xatolik: %s", event.exception)
    for a in await db.fetch("SELECT id FROM admins LIMIT 1"):
        await safe_send(a["id"], f"⚠️ Botda xatolik:\n<code>{esc(str(event.exception))[:500]}</code>")


# ================================================================== ISHGA TUSHIRISH
async def main():
    await db.init()
    for a in ADMIN_IDS:
        await db.run("INSERT OR IGNORE INTO admins(id,full_name) VALUES(?,?)", (a, "Asosiy admin"))
    dp.update.outer_middleware(Activity())
    dp.include_router(r)
    await bot.delete_webhook(drop_pending_updates=True)

    me = await bot.get_me()
    logging.info("Bot ishga tushdi: @%s", me.username)
    reminders = asyncio.create_task(reminder_loop())
    try:
        await dp.start_polling(bot)
    finally:
        reminders.cancel()
        try:
            await reminders
        except asyncio.CancelledError:
            pass


if __name__ == "__main__":
    asyncio.run(main())
